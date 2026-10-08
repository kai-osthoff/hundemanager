# tests/test_hundemanager.py - Sicherung, Update und Rücknahme
#
# Läuft bei jedem Push auf Windows (GitHub Actions) - die App muss IMMER auf
# Windows 11 funktionieren. Lokal:  python -m unittest discover -s tests -v
#
# Jeder Test baut eine eigene Installation in einem Temp-Ordner und spielt
# Updates über einen nachgebauten GitHub-Server ein - echte Releases und der
# echte Dokumente-Ordner werden nie angefasst.

import hashlib
import http.cookiejar
import importlib.util
import io
import json
import locale
import re
import os
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from contextlib import closing
from datetime import date, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Alles, was ausgeliefert wird - bewusst keine handgepflegte Liste, damit neue Dateien
# (z.B. nachweise.py) nicht im Test fehlen
NICHT_AUSGELIEFERT = {'.git', '.github', '.venv', 'instance', 'tests', '__pycache__', '.DS_Store', 'VERSION',
                      '.shepherd-uploads', '.superpowers'}
PROGRAMM = sorted(e for e in os.listdir(REPO) if e not in NICHT_AUSGELIEFERT)
WINDOWS = os.name == 'nt'


# --- Hilfsfunktionen ---

def installation_anlegen(ziel, version):
    os.makedirs(ziel)
    for eintrag in PROGRAMM:
        quelle = os.path.join(REPO, eintrag)
        if os.path.isdir(quelle):
            shutil.copytree(quelle, os.path.join(ziel, eintrag))
        else:
            shutil.copy2(quelle, ziel)
    with open(os.path.join(ziel, 'VERSION'), 'w', encoding='utf-8') as f:
        f.write(version + '\n')
    os.makedirs(os.path.join(ziel, 'instance'))
    return ziel


def db_anlegen(pfad, alt=False):
    """alt=True: Schema wie in v3 (ohne aktiv/fotofreigabe) - testet die Migration mit."""
    person_spalten = '' if alt else ', aktiv BOOLEAN DEFAULT 1, fotofreigabe BOOLEAN DEFAULT 0'
    with closing(sqlite3.connect(pfad)) as v:
        v.executescript(f'''
            CREATE TABLE person (id INTEGER PRIMARY KEY, vorname VARCHAR(100) NOT NULL,
                                 nachname VARCHAR(100) NOT NULL{person_spalten});
            CREATE TABLE hund (id INTEGER PRIMARY KEY, name VARCHAR(100) NOT NULL, geburtstag DATE,
                gueltig_shp_dap_dhp DATE, gueltig_l DATE, gueltig_bbpi DATE, gueltig_t DATE,
                haftpflicht_gueltig BOOLEAN, bemerkung TEXT,
                person_id INTEGER NOT NULL REFERENCES person(id));
            INSERT INTO person (id, vorname, nachname) VALUES
                (1, 'Jürgen', 'Müller'), (2, 'Saskia', 'Test'), (3, 'Zoë', 'Weiß');
            INSERT INTO hund VALUES
                (1, 'Bello', '2020-01-01', '2027-01-01', NULL, NULL, NULL, 1, 'Grüße – „Spezial“', 1),
                (2, 'Luna', NULL, NULL, NULL, NULL, NULL, 0, NULL, 1),
                (3, 'Rex', NULL, NULL, NULL, NULL, NULL, 1, NULL, 2),
                (4, 'Fiete', NULL, NULL, NULL, NULL, NULL, 1, NULL, 3);
        ''')
        v.commit()


def ordner_loeschen(pfad):
    """Wie shutil.rmtree - aber Windows 11 sperrt frische Dateien kurz (Defender-Scan)."""
    for versuch in range(10):
        try:
            shutil.rmtree(pfad)
            return
        except PermissionError:
            if versuch == 9:
                raise
            time.sleep(0.5)


def alter_stand(tag):
    """ZIP eines veröffentlichten Stands. In der Windows-VM gibt es kein git -
    windows-test.sh legt die ZIPs vorher in HUNDEMANAGER_TEST_ARCHIV ab."""
    ordner = os.environ.get('HUNDEMANAGER_TEST_ARCHIV')
    if ordner:
        return lies_bytes(os.path.join(ordner, f'{tag}.zip'))
    return subprocess.run(['git', 'archive', '--format=zip', tag], cwd=REPO,
                          capture_output=True, check=True).stdout


def zaehle(pfad):
    with closing(sqlite3.connect(pfad)) as v:
        return {t: v.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0] for t in ('person', 'hund')}


def lies_bytes(pfad):
    with open(pfad, 'rb') as f:
        return f.read()


def lade_modul(pfad, name):
    spec = importlib.util.spec_from_file_location(name, pfad)
    modul = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modul)
    return modul


def freier_port():
    with closing(socket.socket()) as s:
        s.bind(('127.0.0.1', 0))
        return s.getsockname()[1]


def release_zip(version, aenderungen=None):
    """ZIP wie von GitHub: alles in einem Unterordner. aenderungen: {pfad: inhalt oder None}"""
    puffer = io.BytesIO()
    praefix = 'kai-osthoff-hundemanager-abc1234/'
    aenderungen = dict(aenderungen or {})
    with zipfile.ZipFile(puffer, 'w') as zf:
        for eintrag in PROGRAMM:
            quelle = os.path.join(REPO, eintrag)
            dateien = [quelle] if os.path.isfile(quelle) else [
                os.path.join(w, d) for w, _, ds in os.walk(quelle) for d in ds if not d.endswith('.pyc')]
            for pfad in dateien:
                rel = os.path.relpath(pfad, REPO).replace(os.sep, '/')
                if rel in aenderungen:
                    inhalt = aenderungen.pop(rel)
                    if inhalt is not None:
                        zf.writestr(praefix + rel, inhalt)
                else:
                    zf.write(pfad, praefix + rel)
        zf.writestr(praefix + 'VERSION', version + '\n')
        for rel, inhalt in aenderungen.items():
            if inhalt is not None:
                zf.writestr(praefix + rel, inhalt)
    return puffer.getvalue()


def test_pdf(zeilen):
    """Minimales, gültiges PDF mit Textebene (erfundene Daten - nie echte Dokumente im Repo)."""
    inhalt = 'BT /F1 11 Tf 50 800 Td 14 TL\n' + ''.join(
        '(' + z.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') + ') Tj T*\n' for z in zeilen) + 'ET'
    inhalt = inhalt.encode('cp1252')
    objekte = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R '
        b'/Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length ' + str(len(inhalt)).encode() + b' >>\nstream\n' + inhalt + b'\nendstream',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
    ]
    pdf = bytearray(b'%PDF-1.4\n')
    positionen = []
    for nr, obj in enumerate(objekte, 1):
        positionen.append(len(pdf))
        pdf += f'{nr} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref = len(pdf)
    pdf += f'xref\n0 {len(objekte) + 1}\n0000000000 65535 f \n'.encode()
    for pos in positionen:
        pdf += f'{pos:010d} 00000 n \n'.encode()
    pdf += f'trailer\n<< /Size {len(objekte) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return bytes(pdf)


MONATSNAMEN = ['Januar', 'Februar', 'März', 'April', 'Mai', 'Juni', 'Juli', 'August',
               'September', 'Oktober', 'November', 'Dezember']


def bestaetigung_zeilen(ausgestellt, tier='Bello'):
    """Aufbau wie eine echte Versicherungsbestätigung - mit erfundenen Daten,
    in der durcheinandergewürfelten Reihenfolge, in der pypdf den Text liefert."""
    return [
        'Ihre Tierhalter-Haftpflichtversicherung XY-1234567890 (bitte stets angeben)',
        'Bestaetigung ueber das Bestehen einer Hundehalter-Haftpflichtversicherung',
        'Diese Bestaetigung ist befristet auf die Dauer',
        'eines Jahres ab Ausstellungsdatum.',
        f'1. Tier: {tier}, Labrador; mit Chipnummer 276000000000001; Geburtsdatum: 01.02.2020',
        'Musterversicherung Versicherungs-Aktiengesellschaft',
        'Musterversicherung Versicherungs-AG, 12345 Musterstadt',
        'Max Mustermann',
        'Datum',
        f'{ausgestellt.day:02d}. {MONATSNAMEN[ausgestellt.month - 1]} {ausgestellt.year}',
    ]


def multipart(feldname, dateiname, daten):
    grenze = 'hundemanagertestgrenze'
    koerper = (f'--{grenze}\r\nContent-Disposition: form-data; name="{feldname}"; filename="{dateiname}"\r\n'
               f'Content-Type: application/octet-stream\r\n\r\n').encode() + daten + f'\r\n--{grenze}--\r\n'.encode()
    return koerper, f'multipart/form-data; boundary={grenze}'


def multipart_mehrere(felder, dateien):
    """Formular mit mehreren Dateien: felder = [(name, wert)], dateien = [(name, dateiname, daten)]."""
    grenze = 'hundemanagertestgrenze'
    koerper = b''
    for name, wert in felder:
        koerper += f'--{grenze}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{wert}\r\n'.encode()
    for name, dateiname, daten in dateien:
        koerper += (f'--{grenze}\r\nContent-Disposition: form-data; name="{name}"; filename="{dateiname}"\r\n'
                    f'Content-Type: application/octet-stream\r\n\r\n').encode() + daten + b'\r\n'
    return koerper + f'--{grenze}--\r\n'.encode(), f'multipart/form-data; boundary={grenze}'


# Erfundener Text für Fotos - nie echte Impfpässe in Tests (Datenschutz)
TESTTEXT = ['Impfung gegen Tollwut – gültig bis 09.12.2027', 'Tierarztpraxis Dr. Muster, Hauptstraße 7',
            'Nobivac L4 Lot A450A03 Exp. 11-2026', 'Leptospirose: geimpft am 05.12.2025, gültig bis 05.12.2026',
            'Staupe, Hepatitis, Parvovirose (SHP) – Bello', 'Unterschrift und Stempel des Tierarztes']


def textfoto(art='block', drehung=0, format='JPEG', exif_ausrichtung=None):
    """Foto einer erfundenen Impfpass-Seite, um drehung Grad im Uhrzeigersinn verdreht aufgenommen."""
    from PIL import Image, ImageDraw, ImageFont
    if art == 'block':      # durchgehender Fließtext, linksbündig
        bild = Image.new('RGB', (1400, 1000), (225, 236, 245))
        stift, schrift = ImageDraw.Draw(bild), ImageFont.load_default(size=30)
        for i in range(18):
            stift.text((50, 40 + i * 50), TESTTEXT[i % 6], fill=(30, 30, 60), font=schrift)
    else:                   # verstreute kurze Einträge wie in Tabellenfeldern
        bild = Image.new('RGB', (1000, 1000), 'white')
        stift, schrift = ImageDraw.Draw(bild), ImageFont.load_default(size=24)
        for i in range(12):
            stift.text((40 + (i % 3) * 300, 60 + i * 70), TESTTEXT[i % 6][:18], fill='black', font=schrift)
    bild = bild.rotate(-drehung, expand=True)
    puffer = io.BytesIO()
    if exif_ausrichtung:
        exif = Image.Exif()
        exif[0x0112] = exif_ausrichtung
        bild.save(puffer, format, exif=exif)
    else:
        bild.save(puffer, format)
    return puffer.getvalue()


class GitHubNachbau:
    """Liefert /releases/latest und das ZIP - wie die GitHub-API."""

    def __init__(self, version, zip_daten):
        nachbau = self

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                if self.path == '/latest':
                    daten = json.dumps({
                        'tag_name': 'v' + nachbau.version,
                        'body': 'Testrelease',
                        'zipball_url': f'http://127.0.0.1:{nachbau.port}/zip',
                    }).encode()
                else:
                    daten = nachbau.zip_daten
                self.send_response(200)
                self.send_header('Content-Length', str(len(daten)))
                self.end_headers()
                self.wfile.write(daten)

            def log_message(self, *args):
                pass

        self.version = version
        self.zip_daten = zip_daten
        self.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        self.port = self.server.server_address[1]
        self.url = f'http://127.0.0.1:{self.port}/latest'
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def stop(self):
        self.server.shutdown()
        self.server.server_close()


def prozess_beenden(prozess):
    if prozess.poll() is not None:
        return
    if WINDOWS:
        # start.py startet app.py als Kindprozess - den ganzen Baum beenden
        subprocess.call(['taskkill', '/F', '/T', '/PID', str(prozess.pid)],
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    else:
        import signal
        os.killpg(prozess.pid, signal.SIGKILL)
    prozess.wait(timeout=30)


# --- Sicherungen (direkt im Prozess) ---

class SicherungTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app_dir = installation_anlegen(os.path.join(self.tmp.name, 'Hundemanager'), '5.2.0')
        self.zweiter_ort = os.path.join(self.tmp.name, 'Dokumente', 'Hundemanager-Backups')
        os.environ['HUNDEMANAGER_BACKUP_ZWEITER_ORT'] = self.zweiter_ort
        self.backup = lade_modul(os.path.join(self.app_dir, 'backup.py'), f'backup_{id(self)}')
        db_anlegen(self.backup.DB_PFAD)

    def tearDown(self):
        os.environ.pop('HUNDEMANAGER_BACKUP_ZWEITER_ORT', None)
        self.tmp.cleanup()

    def test_sicherung_ist_vollstaendig_geprueft_und_gespiegelt(self):
        ordner, spiegel_fehler = self.backup.erstelle_backup('manuell', '5.2.0')
        self.assertIsNone(spiegel_fehler)
        manifest = self.backup.pruefe_backup(ordner)
        self.assertEqual(manifest['datensaetze'], {'hund': 4, 'person': 3})
        self.assertIn('code.zip', manifest['dateien'])
        kopie = os.path.join(self.zweiter_ort, os.path.basename(ordner))
        self.backup.pruefe_backup(kopie)
        with zipfile.ZipFile(os.path.join(ordner, 'code.zip')) as zf:
            namen = zf.namelist()
        self.assertIn('app.py', namen)
        self.assertIn('templates/index.html', namen)
        self.assertFalse(any(n.startswith('instance/') for n in namen))

    def test_sicherung_bei_geoeffneter_datenbank(self):
        # Die App hält die Datenbank offen - unter Windows ist die Datei dann gesperrt
        with closing(sqlite3.connect(self.backup.DB_PFAD)) as offen:
            offen.execute("INSERT INTO person (vorname, nachname) VALUES ('Neu', 'Person')")
            offen.commit()
            ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        self.assertEqual(self.backup.pruefe_backup(ordner)['datensaetze']['person'], 4)

    def test_umlaute_bleiben_erhalten(self):
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        with closing(sqlite3.connect(os.path.join(ordner, 'hundemanager.db'))) as v:
            self.assertEqual(v.execute('SELECT bemerkung FROM hund WHERE id=1').fetchone()[0],
                             'Grüße – „Spezial“')

    def test_beschaedigte_sicherung_wird_erkannt(self):
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        with open(os.path.join(ordner, 'hundemanager.db'), 'r+b') as f:
            f.seek(200)
            f.write(b'kaputt')
        with self.assertRaises(self.backup.BackupFehler):
            self.backup.pruefe_backup(ordner)

    def test_unvollstaendige_sicherung_wird_nie_angeboten(self):
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        shutil.copytree(ordner, ordner + '_2' + self.backup.UNVOLLSTAENDIG)
        self.assertEqual([b['name'] for b in self.backup.liste_backups()], [os.path.basename(ordner)])

    def test_sicherung_ohne_manifest_gilt_als_unvollstaendig(self):
        # So sieht eine Sicherung aus, die mittendrin abgebrochen wurde (Strom weg o.ä.)
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        abgebrochen = os.path.join(self.backup.BACKUP_DIR, '2099-01-01_000000_v5.2.0_manuell')
        shutil.copytree(ordner, abgebrochen)
        os.remove(os.path.join(abgebrochen, 'manifest.json'))
        self.assertNotIn(os.path.basename(abgebrochen), [b['name'] for b in self.backup.liste_backups()])
        with self.assertRaises(self.backup.BackupFehler):
            self.backup.finde_backup(os.path.basename(abgebrochen))

    def test_wiederherstellen_sichert_vorher_den_aktuellen_stand(self):
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        with closing(sqlite3.connect(self.backup.DB_PFAD)) as v:
            v.execute('DELETE FROM hund')
            v.commit()
        self.backup.datenbank_wiederherstellen(ordner, '5.2.0')
        self.assertEqual(zaehle(self.backup.DB_PFAD), {'person': 3, 'hund': 4})
        vorher = [b for b in self.backup.liste_backups() if b['art'] == 'vor-wiederherstellung']
        self.assertEqual(len(vorher), 1)
        self.assertEqual(vorher[0]['datensaetze']['hund'], 0)

    def test_wiederherstellen_aus_beschaedigter_sicherung_aendert_nichts(self):
        ordner, _ = self.backup.erstelle_backup('manuell', '5.2.0')
        with open(os.path.join(ordner, 'hundemanager.db'), 'r+b') as f:
            f.seek(200)
            f.write(b'kaputt')
        with closing(sqlite3.connect(self.backup.DB_PFAD)) as v:
            v.execute('DELETE FROM hund WHERE id=4')
            v.commit()
        with self.assertRaises(self.backup.BackupFehler):
            self.backup.datenbank_wiederherstellen(ordner, '5.2.0')
        self.assertEqual(zaehle(self.backup.DB_PFAD), {'person': 3, 'hund': 3})

    def test_aufbewahrung_je_art(self):
        zaehler = iter(range(1000))
        self.backup._zeitstempel = lambda: f'2026-01-01_{next(zaehler):06d}'
        for _ in range(12):
            self.backup.erstelle_backup('vor-update', '5.2.0', mit_code=False)
        for _ in range(3):
            self.backup.erstelle_backup('taeglich', '5.2.0', mit_code=False)
        arten = [b['art'] for b in self.backup.liste_backups()]
        self.assertEqual(arten.count('vor-update'), 10)
        self.assertEqual(arten.count('taeglich'), 3)
        gespiegelt = [n for n in os.listdir(self.zweiter_ort) if n.endswith('vor-update')]
        self.assertEqual(len(gespiegelt), 10)

    def test_sicherung_scheitert_sauber(self):
        # Backup-Ordner nicht anlegbar (z.B. Datei im Weg) -> BackupFehler, kein Absturz
        os.makedirs(self.backup.INSTANCE_DIR, exist_ok=True)
        with open(self.backup.BACKUP_DIR, 'w', encoding='utf-8') as f:
            f.write('im Weg')
        with self.assertRaises(self.backup.BackupFehler):
            self.backup.erstelle_backup('manuell', '5.2.0')


# --- Updates von Ende zu Ende (echte Prozesse wie bei Saskia) ---

class UpdateTests(unittest.TestCase):
    ALT = '5.0.0'
    NEU = '9.9.9'

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app_dir = installation_anlegen(os.path.join(self.tmp.name, 'Hundemanager'), self.ALT)
        self.db = os.path.join(self.app_dir, 'instance', 'hundemanager.db')
        self.zweiter_ort = os.path.join(self.tmp.name, 'Dokumente', 'Hundemanager-Backups')
        self.port = freier_port()
        self.log_pfad = os.path.join(self.tmp.name, 'ausgabe.log')
        self.prozess = None
        self.github = None
        # Wie ein Browser: Cookies behalten, sonst gehen Flask-Meldungen verloren
        self.browser = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.app_py_vorher = lies_bytes(os.path.join(self.app_dir, 'app.py'))
        self.start_bat_vorher = lies_bytes(os.path.join(self.app_dir, 'START.bat'))

    def tearDown(self):
        if self.prozess:
            prozess_beenden(self.prozess)
        if self.github:
            self.github.stop()
        if hasattr(self, '_outcome') and self.log_pfad and os.path.exists(self.log_pfad):
            fehler = [f for f in getattr(self._outcome.result, 'failures', []) +
                      getattr(self._outcome.result, 'errors', []) if f[0] is self]
            if fehler:
                # Die App schreibt in der Kodierung des Systems (Windows: cp1252)
                with open(self.log_pfad, encoding=locale.getpreferredencoding(False), errors='replace') as f:
                    text = f'\n----- Ausgabe von {self.id()} -----\n{f.read()}\n'
                sys.stdout.buffer.write(text.encode('utf-8', 'replace'))
                sys.stdout.flush()
        self.tmp.cleanup()

    def release(self, zip_daten):
        self.github = GitHubNachbau(self.NEU, zip_daten)

    def starten(self, skript='start.py'):
        env = dict(os.environ,
                   HUNDEMANAGER_PORT=str(self.port),
                   HUNDEMANAGER_RELEASES_URL=self.github.url,
                   HUNDEMANAGER_BACKUP_ZWEITER_ORT=self.zweiter_ort,
                   HUNDEMANAGER_KEIN_BROWSER='1', **getattr(self, 'extra_env', {}))
        env.pop('HUNDEMANAGER_LAUNCHER', None)
        optionen = {'creationflags': subprocess.CREATE_NEW_PROCESS_GROUP} if WINDOWS else {'start_new_session': True}
        self.log = open(self.log_pfad, 'w', encoding='utf-8')
        self.addCleanup(self.log.close)
        self.prozess = subprocess.Popen([sys.executable, skript], cwd=self.app_dir, env=env,
                                        stdout=self.log, stderr=subprocess.STDOUT, **optionen)
        self.assertEqual(self.warte_auf_version(), self.ALT)

    def url(self, pfad):
        return f'http://127.0.0.1:{self.port}{pfad}'

    def status_abfragen(self):
        with urllib.request.urlopen(self.url('/api/version'), timeout=5) as a:
            return json.load(a)

    def version_abfragen(self):
        return self.status_abfragen()['version']

    def warte_auf_version(self, timeout=90):
        ende = time.time() + timeout
        while time.time() < ende:
            try:
                status = self.status_abfragen()
                self.instanz = status.get('instanz')  # v5.1.0 kennt noch keine Kennung
                return status['version']
            except (urllib.error.URLError, ConnectionError, OSError):
                time.sleep(0.5)
        self.fail('Hundemanager antwortet nicht')

    def warte_auf_neustart(self, timeout=90):
        """Wartet, bis ein NEUER Prozess antwortet (andere Kennung). Gibt dessen Version zurück.

        Nicht an "Server war kurz weg" festmachen: unter Windows warten Verbindungs-
        versuche ~2 s und landen nahtlos beim neuen Prozess - der Server ist nie "weg".
        """
        ende = time.time() + timeout
        while time.time() < ende:
            try:
                status = self.status_abfragen()
                if status.get('instanz') and status.get('instanz') != self.instanz:
                    return status['version']
            except (urllib.error.URLError, ConnectionError, OSError):
                pass
            time.sleep(0.5)
        self.fail('Kein Neustart nach dem Update')

    def update_klicken(self):
        anfrage = urllib.request.Request(self.url('/update/installieren'), data=b'', method='POST')
        with self.browser.open(anfrage, timeout=600) as a:
            return a.read().decode('utf-8')

    def seite(self, pfad='/'):
        with self.browser.open(self.url(pfad), timeout=10) as a:
            return a.read().decode('utf-8')

    def sicherungen(self, art):
        basis = os.path.join(self.app_dir, 'instance', 'backup')
        return [n for n in os.listdir(basis) if n.endswith('_' + art)]

    def pruefe_daten_unveraendert(self):
        self.assertEqual(zaehle(self.db), {'person': 3, 'hund': 4})
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT nachname FROM person WHERE id=1').fetchone()[0], 'Müller')

    def pruefe_alte_version_wiederhergestellt(self):
        self.assertEqual(lies_bytes(os.path.join(self.app_dir, 'app.py')), self.app_py_vorher)
        with open(os.path.join(self.app_dir, 'VERSION'), encoding='utf-8') as f:
            self.assertEqual(f.read().strip(), self.ALT)
        self.assertFalse(os.path.exists(os.path.join(self.app_dir, 'instance', 'update_laeuft.json')))

    def test_update_mit_geprueftem_backup(self):
        db_anlegen(self.db, alt=True)  # alte Datenbank: Migration läuft beim Update mit
        self.release(release_zip(self.NEU))
        self.starten()

        kopf = self.seite()
        self.assertIn('Neue Version 9.9.9 installieren', kopf)
        self.assertIn('Testrelease', kopf)
        self.assertNotIn('Updates suchen', kopf)
        self.assertIn('wurde installiert', self.update_klicken())
        self.assertEqual(self.warte_auf_neustart(), self.NEU)

        self.pruefe_daten_unveraendert()
        sicherungen = self.sicherungen('vor-update')
        self.assertEqual(len(sicherungen), 1)
        self.assertIn('_v5.0.0_', sicherungen[0])
        backup = lade_modul(os.path.join(self.app_dir, 'backup.py'), f'backup_{id(self)}')
        for basis in (os.path.join(self.app_dir, 'instance', 'backup'), self.zweiter_ort):
            manifest = backup.pruefe_backup(os.path.join(basis, sicherungen[0]))
            self.assertEqual(manifest['datensaetze']['hund'], 4)
            self.assertEqual(manifest['datensaetze']['person'], 3)
        self.assertFalse(os.path.exists(os.path.join(self.app_dir, 'instance', 'update_laeuft.json')))
        self.assertEqual(lies_bytes(os.path.join(self.app_dir, 'START.bat')), self.start_bat_vorher)
        self.assertNotIn('hat nicht geklappt', self.seite())

    def test_ohne_neue_version_steht_oben_updates_suchen(self):
        self.github = GitHubNachbau(self.ALT, release_zip(self.ALT))
        self.starten()

        for pfad in ('/', '/nachweise', '/sicherungen'):
            seite = self.seite(pfad)
            self.assertIn('Updates suchen', seite)
            self.assertNotIn('installieren</button>', seite)
        anfrage = urllib.request.Request(self.url('/update/pruefen'), data=b'', method='POST')
        with self.browser.open(anfrage, timeout=30) as a:
            self.assertIn('auf dem neuesten Stand', a.read().decode('utf-8'))

    def test_neues_release_erscheint_ohne_neustart(self):
        """Ein Release, das erst kommt, während die App läuft, taucht oben rechts auf."""
        self.github = GitHubNachbau(self.ALT, release_zip(self.NEU))
        self.extra_env = {'HUNDEMANAGER_PRUEF_INTERVALL': '1'}
        self.starten()
        self.assertIn('Updates suchen', self.seite('/update/kopf'))
        self.assertIn("fetch('/update/kopf'", self.seite())  # offene Seite fragt nach

        self.github.version = self.NEU
        ende = time.time() + 30
        while 'Neue Version 9.9.9 installieren' not in self.seite('/update/kopf'):
            if time.time() > ende:
                self.fail('Neues Release wurde nicht erkannt')
            time.sleep(0.5)
        self.assertIn('Neue Version 9.9.9 installieren', self.seite())

    def test_kaputtes_programm_wird_zurueckgenommen(self):
        db_anlegen(self.db)
        with open(os.path.join(REPO, 'app.py'), encoding='utf-8') as f:
            kaputt = f.read() + '\ndas ist kein python (\n'
        self.release(release_zip(self.NEU, {'app.py': kaputt, 'neue_datei.py': 'x = 1\n'}))
        self.starten()

        self.assertIn('wurde installiert', self.update_klicken())
        self.assertEqual(self.warte_auf_neustart(), self.ALT)

        self.pruefe_alte_version_wiederhergestellt()
        self.pruefe_daten_unveraendert()
        self.assertFalse(os.path.exists(os.path.join(self.app_dir, 'neue_datei.py')))
        self.assertIn('hat nicht geklappt', self.seite())

    def test_selbsttest_fehler_wird_zurueckgenommen(self):
        db_anlegen(self.db)
        self.release(release_zip(self.NEU, {'templates/index.html': '{% extends "base.html" %}{% block content %}{{ kaputt( }}{% endblock %}'}))
        self.starten()

        self.update_klicken()
        self.assertEqual(self.warte_auf_neustart(), self.ALT)

        self.pruefe_alte_version_wiederhergestellt()
        self.pruefe_daten_unveraendert()
        self.assertIn('hat nicht geklappt', self.seite())

    def test_kopierfehler_wird_sofort_zurueckgenommen(self):
        db_anlegen(self.db)
        # start.py ist ein Ordner -> das Überschreiben schlägt mitten im Update fehl
        os.remove(os.path.join(self.app_dir, 'start.py'))
        os.makedirs(os.path.join(self.app_dir, 'start.py'))
        self.app_py_vorher = lies_bytes(os.path.join(self.app_dir, 'app.py'))
        self.release(release_zip(self.NEU, {'neue_datei.py': 'x = 1\n'}))
        self.starten('app.py')

        antwort = self.update_klicken()
        self.assertIn('wurde wiederhergestellt', antwort)
        self.assertEqual(self.version_abfragen(), self.ALT)
        self.pruefe_alte_version_wiederhergestellt()
        self.pruefe_daten_unveraendert()
        self.assertFalse(os.path.exists(os.path.join(self.app_dir, 'neue_datei.py')))

    def test_kein_update_ohne_backup(self):
        db_anlegen(self.db)
        self.release(release_zip(self.NEU))
        self.starten()
        # Ab jetzt kann keine Sicherung mehr angelegt werden
        backup_dir = os.path.join(self.app_dir, 'instance', 'backup')
        ordner_loeschen(backup_dir)
        with open(backup_dir, 'w', encoding='utf-8') as f:
            f.write('im Weg')

        antwort = self.update_klicken()
        self.assertIn('Sicherung nicht geklappt', antwort)
        self.assertEqual(self.version_abfragen(), self.ALT)
        self.pruefe_alte_version_wiederhergestellt()
        self.pruefe_daten_unveraendert()

    def test_update_von_5_1_2_uebernimmt_haftpflicht_nachweise(self):
        """5.1.2 hatte eine eigene Haftpflicht-Tabelle - nach dem Update muss alles in der
        gemeinsamen Nachweis-Ablage stehen, die alte Tabelle bleibt unangetastet."""
        ordner_loeschen(self.app_dir)
        os.makedirs(self.app_dir)
        with zipfile.ZipFile(io.BytesIO(alter_stand('v5.1.2'))) as zf:
            zf.extractall(self.app_dir)  # git archive: ohne Oberordner
        os.makedirs(os.path.join(self.app_dir, 'instance'), exist_ok=True)
        db_anlegen(self.db)
        self.release(release_zip(self.NEU))
        self.ALT = '5.1.2'
        self.starten()

        # Wie Saskia in 5.1.2: Nachweis über die alte Oberfläche hochladen
        pdf = test_pdf(bestaetigung_zeilen(date.today()))
        koerper, typ = multipart('datei', 'Bestätigung.pdf', pdf)
        anfrage = urllib.request.Request(self.url('/hund/1/haftpflicht/hochladen'), data=koerper, method='POST')
        anfrage.add_header('Content-Type', typ)
        seite = self.browser.open(anfrage, timeout=30).read().decode('utf-8')
        sha = re.search(r'name="sha256" value="([0-9a-f]{64})"', seite).group(1)
        formular = urllib.parse.urlencode({
            'sha256': sha, 'endung': 'pdf', 'original_name': 'Bestätigung.pdf',
            'gueltig_bis': '2027-10-04', 'ausgestellt_am': '2026-10-05',
            'versicherer': 'Musterversicherung AG', 'vertragsnummer': 'XY-1234567890',
            'tier_laut_nachweis': 'Bello', 'chipnummer': '276000000000001',
        }).encode()
        self.browser.open(urllib.request.Request(self.url('/hund/1/haftpflicht/speichern'),
                                                 data=formular, method='POST'), timeout=30).read()

        self.assertIn('wurde installiert', self.update_klicken())
        self.assertEqual(self.warte_auf_neustart(), self.NEU)

        historie = self.seite('/nachweise')
        for erwartet in ('XY-1234567890', 'Musterversicherung AG', 'Jürgen Müller', 'Bello'):
            self.assertIn(erwartet, historie)
        self.assertIn('Haftpflicht bis 04.10.2027', self.seite('/'))
        self.assertEqual(self.browser.open(self.url('/nachweis/1/datei'), timeout=10).read(), pdf)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT COUNT(*) FROM haftpflicht_nachweis').fetchone()[0], 1)
            self.assertEqual(v.execute("SELECT COUNT(*) FROM nachweis WHERE art='haftpflicht'").fetchone()[0], 1)
        self.assertGreaterEqual(len(self.sicherungen('vor-migration')), 1)

    def test_wiederherstellen_ueber_die_oberflaeche(self):
        db_anlegen(self.db)
        self.release(release_zip(self.NEU))
        self.starten()
        name = self.sicherungen('taeglich')[0]  # beim Start automatisch angelegt
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('DELETE FROM hund')
            v.commit()

        anfrage = urllib.request.Request(self.url('/sicherungen/wiederherstellen'),
                                         data=f'name={name}'.encode(), method='POST')
        with self.browser.open(anfrage, timeout=30) as a:
            self.assertIn('wiederhergestellt', a.read().decode('utf-8'))
        self.pruefe_daten_unveraendert()
        self.assertIn('Bello', self.seite())
        self.assertEqual(len(self.sicherungen('vor-wiederherstellung')), 1)


    def test_update_von_saskias_5_1_0(self):
        """Der echte Übergang bei Saskia: alter Updater und alter Starter aus v5.1.0."""
        ordner_loeschen(self.app_dir)
        os.makedirs(self.app_dir)
        archiv = alter_stand('v5.1.0')
        with zipfile.ZipFile(io.BytesIO(archiv)) as zf:
            zf.extractall(self.app_dir)
        os.makedirs(os.path.join(self.app_dir, 'instance'))
        db_anlegen(self.db)
        self.release(release_zip(self.NEU))
        # v5.1.0 kennt keine Umgebungsvariable für die Release-Adresse - nur für den Test umbiegen
        updater_py = os.path.join(self.app_dir, 'updater.py')
        with open(updater_py, encoding='utf-8') as f:
            inhalt = f.read()
        with open(updater_py, 'w', encoding='utf-8') as f:
            f.write(inhalt.replace(
                "f'https://api.github.com/repos/{GITHUB_REPO}/releases/latest'", repr(self.github.url)))
        self.ALT = '5.1.0'
        self.starten()

        self.update_klicken()
        self.assertEqual(self.warte_auf_neustart(), self.NEU)
        self.pruefe_daten_unveraendert()
        # Der alte Starter läuft noch -> Hinweis, einmal neu zu starten
        self.assertIn('START.bat', self.seite())
        self.assertIn('neu öffnen', self.seite())
        # Die neue Version legt sofort ihre erste geprüfte Sicherung an
        self.assertEqual(len(self.sicherungen('taeglich')), 1)


# --- Haftpflicht-Nachweise ---

class NachweisTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app_dir = installation_anlegen(os.path.join(self.tmp.name, 'Hundemanager'), '5.1.2')
        self.zweiter_ort = os.path.join(self.tmp.name, 'Dokumente', 'Hundemanager-Backups')
        os.environ['HUNDEMANAGER_BACKUP_ZWEITER_ORT'] = self.zweiter_ort
        # nachweise.py importiert backup - beide aus der Testinstallation laden
        sys.path.insert(0, self.app_dir)
        for name in ('backup', 'nachweise'):
            sys.modules.pop(name, None)
        import backup, nachweise
        self.backup, self.nachweise = backup, nachweise
        db_anlegen(backup.DB_PFAD)

    def tearDown(self):
        sys.path.remove(self.app_dir)
        for name in ('backup', 'nachweise'):
            sys.modules.pop(name, None)
        os.environ.pop('HUNDEMANAGER_BACKUP_ZWEITER_ORT', None)
        self.tmp.cleanup()

    def test_angaben_werden_erkannt(self):
        pdf = test_pdf(bestaetigung_zeilen(date(2026, 10, 5)))
        vorschlag = self.nachweise.angaben_vorschlagen(self.nachweise.pdf_text(pdf))
        self.assertEqual(vorschlag['versicherer'], 'Musterversicherung Versicherungs-Aktiengesellschaft')
        self.assertEqual(vorschlag['vertragsnummer'], 'XY-1234567890')
        self.assertEqual(vorschlag['ausgestellt_am'], date(2026, 10, 5))  # nicht das Geburtsdatum
        self.assertEqual(vorschlag['gueltig_bis'], date(2027, 10, 4))     # ein Jahr ab Ausstellung
        self.assertEqual(vorschlag['tier'], 'Bello')
        self.assertEqual(vorschlag['chipnummer'], '276000000000001')

    def test_ausdrueckliches_gueltig_bis(self):
        text = 'Muster Versicherung AG\nDatum 01.03.2026\nDer Versicherungsschutz ist gültig bis 31.12.2026.'
        self.assertEqual(self.nachweise.angaben_vorschlagen(text)['gueltig_bis'], date(2026, 12, 31))

    def test_ohne_text_keine_vorschlaege(self):
        self.assertEqual(self.nachweise.angaben_vorschlagen(''), {})
        self.assertEqual(self.nachweise.pdf_text(b'%PDF-kaputt'), '')

    def test_speichern_ist_unveraenderlich_und_geprueft(self):
        pdf = test_pdf(['Test'])
        sha, endung = self.nachweise.speichern(pdf)
        self.assertEqual((sha, endung), self.nachweise.speichern(pdf))  # gleicher Inhalt, gleiche Datei
        self.assertEqual(lies_bytes(self.nachweise.pfad(sha, endung)), pdf)
        self.assertEqual(os.listdir(self.nachweise.NACHWEIS_DIR), [f'{sha}.pdf'])
        self.assertEqual(self.nachweise.speichern(b'\xff\xd8\xff\xe0foto')[1], 'jpg')
        for falsch in (b'', b'MZ ausfuehrbar', b'<html>'):
            with self.assertRaises(self.nachweise.NachweisFehler):
                self.nachweise.speichern(falsch)
        with self.assertRaises(self.nachweise.NachweisFehler):
            self.nachweise.pfad('../../app', 'pdf')

    def test_nachweise_sind_in_jeder_sicherung(self):
        sha, endung = self.nachweise.speichern(test_pdf(['Nachweis A']))
        name = f'{sha}.{endung}'
        ordner, hinweis = self.backup.erstelle_backup('manuell', '5.1.2')
        self.assertIsNone(hinweis)
        self.assertEqual(self.backup.pruefe_backup(ordner)['nachweise'], [name])
        for basis in (self.backup.BACKUP_DIR, self.zweiter_ort):
            self.assertTrue(os.path.exists(os.path.join(basis, 'nachweise', name)))
        # Eine zweite Sicherung legt die Datei nicht noch einmal ab
        self.backup.erstelle_backup('manuell', '5.1.2')
        self.assertEqual(os.listdir(os.path.join(self.backup.BACKUP_DIR, 'nachweise')), [name])

    def test_verlorener_nachweis_wird_wiederhergestellt(self):
        sha, endung = self.nachweise.speichern(test_pdf(['Nachweis B']))
        ordner, _ = self.backup.erstelle_backup('manuell', '5.1.2')
        os.remove(self.nachweise.pfad(sha, endung))
        self.backup.datenbank_wiederherstellen(ordner, '5.1.2')
        self.assertEqual(self.backup._sha256(self.nachweise.pfad(sha, endung)), sha)

    def test_beschaedigter_nachweis_wird_gemeldet_und_nicht_ueberschrieben(self):
        sha, endung = self.nachweise.speichern(test_pdf(['Nachweis C']))
        self.backup.erstelle_backup('manuell', '5.1.2')
        with open(self.nachweise.pfad(sha, endung), 'r+b') as f:
            f.write(b'XXXX')
        ordner, hinweis = self.backup.erstelle_backup('manuell', '5.1.2')
        self.assertIn('beschädigt', hinweis)
        # Die intakte Kopie in den Sicherungen bleibt intakt
        self.backup.pruefe_backup(ordner)

    def test_beschaedigte_ablage_macht_sicherung_ungueltig(self):
        sha, endung = self.nachweise.speichern(test_pdf(['Nachweis D']))
        ordner, _ = self.backup.erstelle_backup('manuell', '5.1.2')
        with open(os.path.join(self.backup.BACKUP_DIR, 'nachweise', f'{sha}.{endung}'), 'r+b') as f:
            f.write(b'XXXX')
        with self.assertRaises(self.backup.BackupFehler):
            self.backup.pruefe_backup(ordner)


class HaftpflichtOberflaecheTests(unittest.TestCase):
    """Hochladen, Prüfen, Speichern und Ansehen über die echte Oberfläche."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.app_dir = installation_anlegen(os.path.join(self.tmp.name, 'Hundemanager'), '5.1.2')
        self.db = os.path.join(self.app_dir, 'instance', 'hundemanager.db')
        db_anlegen(self.db)
        if hasattr(self, 'vor_dem_start'):
            self.vor_dem_start()  # z.B. Datenbank im Stand einer älteren Version
        self.port = freier_port()
        self.log_pfad = os.path.join(self.tmp.name, 'ausgabe.log')
        env = dict(os.environ, HUNDEMANAGER_PORT=str(self.port), HUNDEMANAGER_KEIN_BROWSER='1',
                   HUNDEMANAGER_RELEASES_URL='http://127.0.0.1:9/gibt-es-nicht',
                   HUNDEMANAGER_BACKUP_ZWEITER_ORT=os.path.join(self.tmp.name, 'Dokumente'))
        env.pop('HUNDEMANAGER_LAUNCHER', None)
        optionen = {'creationflags': subprocess.CREATE_NEW_PROCESS_GROUP} if WINDOWS else {'start_new_session': True}
        self.log = open(self.log_pfad, 'w', encoding='utf-8')
        self.prozess = subprocess.Popen([sys.executable, 'app.py'], cwd=self.app_dir, env=env,
                                        stdout=self.log, stderr=subprocess.STDOUT, **optionen)
        self.browser = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        ende = time.time() + 60
        while time.time() < ende:
            try:
                self.browser.open(f'http://127.0.0.1:{self.port}/api/version', timeout=5).read()
                return
            except OSError:
                time.sleep(0.5)
        self.fail('Hundemanager antwortet nicht')

    def tearDown(self):
        prozess_beenden(self.prozess)
        self.log.close()
        self.tmp.cleanup()

    def oeffne(self, pfad, daten=None, typ=None):
        anfrage = urllib.request.Request(f'http://127.0.0.1:{self.port}{pfad}', data=daten,
                                         method='POST' if daten is not None else 'GET')
        if typ:
            anfrage.add_header('Content-Type', typ)
        with self.browser.open(anfrage, timeout=30) as a:
            return a.read()

    def test_nachweis_hochladen_pruefen_speichern_ansehen(self):
        heute = date.today()
        pdf = test_pdf(bestaetigung_zeilen(heute))
        koerper, typ = multipart('datei', 'Bestätigung Haftpflicht.pdf', pdf)

        seite = self.oeffne('/hund/1/haftpflicht/hochladen', koerper, typ).decode('utf-8')
        self.assertIn('Angaben wurden aus dem Dokument erkannt', seite)
        self.assertIn('XY-1234567890', seite)
        gueltig_bis = heute.replace(year=heute.year + 1) - timedelta(days=1) if not (heute.month == 2 and heute.day == 29) \
            else date(heute.year + 1, 2, 28)
        self.assertIn(gueltig_bis.isoformat(), seite)
        sha = re.search(r'name="sha256" value="([0-9a-f]{64})"', seite).group(1)

        formular = urllib.parse.urlencode({
            'sha256': sha, 'endung': 'pdf', 'original_name': 'Bestätigung Haftpflicht.pdf',
            'gueltig_bis': gueltig_bis.isoformat(), 'ausgestellt_am': heute.isoformat(),
            'versicherer': 'Musterversicherung', 'vertragsnummer': 'XY-1234567890',
            'tier_laut_nachweis': 'Bello', 'chipnummer': '276000000000001',
        }).encode()
        seite = self.oeffne('/hund/1/haftpflicht/speichern', formular).decode('utf-8')
        self.assertIn('Haftpflicht-Nachweis für Bello gespeichert', seite)

        uebersicht = self.oeffne('/').decode('utf-8')
        self.assertIn(f'Haftpflicht bis {gueltig_bis.strftime("%d.%m.%Y")}', uebersicht)
        self.assertIn('Nachweis ansehen', uebersicht)
        self.assertEqual(self.oeffne('/nachweis/1/datei'), pdf)

        # Excel-Export zeigt die Gültigkeit
        import openpyxl
        mappe = openpyxl.load_workbook(io.BytesIO(self.oeffne('/export/excel')))
        werte = [z[9] for z in mappe.active.iter_rows(min_row=2, values_only=True)]
        self.assertIn(f'bis {gueltig_bis.strftime("%d.%m.%Y")}', werte)

        # Falscher Hund im Dokument -> Warnung
        koerper, typ = multipart('datei', 'anderer.pdf', test_pdf(bestaetigung_zeilen(heute, tier='Rocky')))
        self.assertIn('Rocky', self.oeffne('/hund/1/haftpflicht/hochladen', koerper, typ).decode('utf-8'))

    def test_dateien_hineinziehen(self):
        # Übersicht: Hund und Halter sind Ablageziele, ein Dialog fragt, was es ist
        uebersicht = self.oeffne('/').decode('utf-8')
        self.assertIn('einfach auf einen Hund oder Halter ziehen', uebersicht)
        self.assertIn('id="ablage-dialog"', uebersicht)
        self.assertIn('data-ablage-hund="Bello" data-ablage-url="/hund/1/haftpflicht/hochladen"', uebersicht)
        self.assertIn('data-ablage-person="Jürgen Müller" data-ablage-url="/person/1/fotoeinwilligung/hochladen"',
                      uebersicht)
        self.assertIn('{"haftpflicht": "Haftpflicht", "impfpass": "Impfpass"}', uebersicht)
        # Upload-Seiten: Ablagefläche; beim Nachweis wird gleich hochgeladen, bei der Einwilligung nicht
        self.assertIn('<label class="ablage" data-senden>', self.oeffne('/hund/1/nachweise?art=impfpass').decode('utf-8'))
        self.assertIn('<label class="ablage">', self.oeffne('/person/1/fotoeinwilligung').decode('utf-8'))

        # So schickt der Dialog ab: Datei an die gewählte Adresse, die Art kommt vom Knopf
        koerper, typ = multipart_mehrere([('art', 'impfpass')], [('datei', 'impfpass.jpg', textfoto('felder'))])
        self.assertIn('Impfpass-Nachweis eintragen', self.oeffne('/hund/1/nachweise/hochladen', koerper, typ).decode('utf-8'))
        koerper, typ = multipart_mehrere([('art', 'fotoeinwilligung')], [('datei', 'einwilligung.pdf', test_pdf(['Ja']))])
        self.assertIn('Fotoeinwilligung von Jürgen Müller gespeichert',
                      self.oeffne('/person/1/fotoeinwilligung/hochladen', koerper, typ).decode('utf-8'))

    def test_uebersicht_filter_und_pausierte(self):
        seite = self.oeffne('/').decode('utf-8')
        for name in ('Bello', 'Luna', 'Rex', 'Fiete'):
            self.assertIn(name, seite)
        # Ohne Impfdaten ist bei allen etwas zu tun
        self.assertIn('Handlungsbedarf (4)', seite)
        self.assertIn('Rex', self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8'))
        self.assertIn('Bello', self.oeffne('/?all=1').decode('utf-8'))  # alte Links gehen weiter

        # Der Fall aus der Praxis: alle pausiert -> nie "keine Einträge", sondern ein Weg zurück
        for person_id in (1, 2, 3):
            self.oeffne(f'/person/{person_id}/toggle_aktiv', b'')
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('Gerade ist niemand aktiv', seite)
        self.assertIn('Pausiert (3)', seite)
        self.assertNotIn('Noch keine Einträge', seite)
        self.assertIn('Bello', self.oeffne('/?ansicht=alle').decode('utf-8'))

    def test_hundekarte_und_impfungen_klickbar(self):
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE hund SET gueltig_t = ? WHERE id = 1', (date(2027, 5, 14).isoformat(),))
            v.commit()
        seite = self.oeffne('/').decode('utf-8')
        # Die ganze Karte führt zu den Nachweisen, jede Impfung zum Impfpass
        self.assertIn('<a href="/hund/1/nachweise" class="karte-link">Bello</a>', seite)
        self.assertEqual(len(re.findall(r'<a href="/hund/1/nachweise\?art=impfpass" class="kachel ', seite)), 4)
        self.assertIn('<strong>05/27</strong>', seite)
        self.assertIn('T: gültig bis 14.05.2027', seite)
        # Halter-Aktionen stecken in Menüs statt in einer Knopfreihe
        self.assertIn('class="menue', seite)
        self.assertIn('Pausieren</button>', seite)

    def test_kennzahl_kacheln_filtern(self):
        heute = date.today()
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE hund SET gueltig_l = ? WHERE name = ?', ((heute - timedelta(days=1)).isoformat(), 'Rex'))
            v.execute('UPDATE hund SET gueltig_t = ? WHERE name = ?', ((heute + timedelta(days=10)).isoformat(), 'Luna'))
            v.commit()

        def hunde(seite):
            return {n for n in ('Bello', 'Luna', 'Rex', 'Fiete') if f'aria-label="{n}"' in seite}

        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('href="/?filter=abgelaufen"', seite)
        self.assertEqual(hunde(seite), {'Bello', 'Luna', 'Rex', 'Fiete'})

        seite = self.oeffne('/?filter=abgelaufen').decode('utf-8')
        self.assertEqual(hunde(seite), {'Rex'})
        # Aktive Kachel führt zurück zur ungefilterten Liste
        self.assertRegex(seite, r'<a href="/"\s+class="kennzahl[^"]*" aria-current="true" title="Filter aufheben"')
        self.assertEqual(hunde(self.oeffne('/?filter=bald').decode('utf-8')), {'Luna'})
        self.assertEqual(hunde(self.oeffne('/?filter=haftpflicht').decode('utf-8')), {'Bello', 'Luna', 'Rex', 'Fiete'})
        self.assertEqual(hunde(self.oeffne('/?filter=foto').decode('utf-8')), {'Bello', 'Luna', 'Rex', 'Fiete'})
        self.assertEqual(hunde(self.oeffne('/?filter=quatsch').decode('utf-8')), {'Bello', 'Luna', 'Rex', 'Fiete'})

        # Pausierte Halter fallen raus, wie bei den Kennzahlen
        self.oeffne('/person/2/toggle_aktiv', b'')
        seite = self.oeffne('/?filter=abgelaufen').decode('utf-8')
        self.assertEqual(hunde(seite), set())
        self.assertIn('Nichts gefunden', seite)

    def nachweis_speichern(self, art, datei, dateiname, felder):
        koerper, typ = multipart('datei', dateiname, datei)
        # Art steckt im Upload-Formular als eigenes Feld
        grenze = typ.split('boundary=')[1]
        koerper = (f'--{grenze}\r\nContent-Disposition: form-data; name="art"\r\n\r\n{art}\r\n').encode() + koerper
        seite = self.oeffne('/hund/1/nachweise/hochladen', koerper, typ).decode('utf-8')
        sha = re.search(r'name="sha256" value="([0-9a-f]{64})"', seite).group(1)
        endung = re.search(r'name="endung" value="(\w+)"', seite).group(1)
        daten = [('art', art), ('sha256', sha), ('endung', endung), ('original_name', dateiname)]
        for name, wert in felder.items():
            for w in (wert if isinstance(wert, list) else [wert]):
                daten.append((name, w))
        return self.oeffne('/hund/1/nachweise/speichern', urllib.parse.urlencode(daten).encode()).decode('utf-8')

    def test_impfpass_foto_in_derselben_ablage(self):
        foto = b'\xff\xd8\xff\xe0' + b'Impfpass-Foto' * 50
        seite = self.nachweis_speichern('impfpass', foto, 'impfpass.jpg', {
            'eingereicht_von': 'Jürgen Müller', 'aussteller': 'Tierarztpraxis Muster',
            'gueltig_bis': '2027-03-01', 'impfungen': ['SHP/DAP/DHP', 'L']})
        self.assertIn('Impfpass-Nachweis für Bello gespeichert', seite)
        self.assertIn('SHP/DAP/DHP, L', seite)
        historie = self.oeffne('/nachweise?art=impfpass').decode('utf-8')
        self.assertIn('Tierarztpraxis Muster', historie)
        self.assertEqual(self.oeffne('/nachweis/1/datei'), foto)
        # Impfpass ohne Gültigkeitsdatum ist erlaubt, Haftpflicht nicht
        seite = self.nachweis_speichern('impfpass', foto + b'2', 'seite2.jpg', {'impfungen': ['T']})
        self.assertIn('Impfpass-Nachweis für Bello gespeichert', seite)
        seite = self.nachweis_speichern('haftpflicht', foto + b'3', 'haft.jpg', {'gueltig_bis': ''})
        self.assertIn('Bitte angeben, bis wann der Nachweis gültig ist', seite)

    def test_bald_ablaufende_nachweise_zum_erinnern(self):
        bald = (date.today() + timedelta(days=20)).isoformat()
        self.nachweis_speichern('haftpflicht', test_pdf(['A']), 'alt.pdf', {'gueltig_bis': bald})
        seite = self.oeffne('/nachweise').decode('utf-8')
        self.assertIn('Neuen hochladen', seite)
        self.assertIn('in 20 Tagen', seite)
        # Neuer Nachweis eingereicht -> keine Erinnerung mehr, alter bleibt in der Historie
        spaeter = (date.today() + timedelta(days=365)).isoformat()
        self.nachweis_speichern('haftpflicht', test_pdf(['B']), 'neu.pdf', {'gueltig_bis': spaeter})
        seite = self.oeffne('/nachweise').decode('utf-8')
        self.assertIn('Gerade muss niemand einen neuen Nachweis einreichen', seite)
        self.assertIn('alt.pdf', self.oeffne('/hund/1/nachweise').decode('utf-8'))
        self.assertIn('neu.pdf', self.oeffne('/hund/1/nachweise').decode('utf-8'))

    def test_falsche_datei_wird_abgelehnt(self):
        koerper, typ = multipart('datei', 'virus.exe', b'MZ\x90\x00')
        seite = self.oeffne('/hund/1/haftpflicht/hochladen', koerper, typ).decode('utf-8')
        self.assertIn('Bitte ein PDF oder ein Foto', seite)
        self.assertFalse(os.path.exists(os.path.join(self.app_dir, 'instance', 'nachweise')))



class WhatsAppTests(unittest.TestCase):
    """Click-to-Chat-Links: Nummer international, Text URL-codiert, kein externer Dienst."""

    def setUp(self):
        self.wa = lade_modul(os.path.join(REPO, 'whatsapp.py'), 'whatsapp_test')

    def test_uebliche_schreibweisen(self):
        for eingabe in ('0171 1234567', '0171/123 45 67', '0171-1234567', '+49 171 1234567',
                        '0049 171 1234567', '+49 (0) 171 1234567', '(0171) 1234567'):
            self.assertEqual(self.wa.nummer(eingabe), '491711234567', eingabe)
        self.assertEqual(self.wa.nummer('+43 664 1234567'), '436641234567')
        self.assertEqual(self.wa.nummer(' 0151.23456789 '), '4915123456789')

    def test_unsinnige_nummern(self):
        for eingabe in ('', None, '   ', 'keine', '0171 12a4567', '123', '+49 171 1234567 1234567',
                        '+0171 1234567', '00 0171 1234567'):
            self.assertIsNone(self.wa.nummer(eingabe), eingabe)

    def test_link_ohne_und_mit_text(self):
        self.assertEqual(self.wa.link('0171 1234567'),
                         'https://web.whatsapp.com/send?phone=491711234567')
        self.assertIsNone(self.wa.link(''))
        link = self.wa.link('0171 1234567', 'Hallo Zoë,\nImpfung SHP/DAP/DHP & L: 50 %?')
        abfrage = urllib.parse.urlparse(link).query
        self.assertNotIn(' ', link)
        self.assertNotIn('+', abfrage)  # Leerzeichen als %20, nicht als +
        self.assertEqual(urllib.parse.parse_qs(abfrage)['text'], ['Hallo Zoë,\nImpfung SHP/DAP/DHP & L: 50 %?'])
        self.assertEqual(urllib.parse.parse_qs(abfrage)['phone'], ['491711234567'])

    def test_fotoeinwilligung_nachricht_mit_link(self):
        link = 'https://example.org/formular.pdf'
        text = self.wa.fotoeinwilligung_nachricht('Zoë', ['Bello', 'Luna', 'Rex'], link)
        self.assertTrue(text.startswith('Hallo Zoë,'))
        self.assertIn('mit Bello, Luna und Rex', text)
        self.assertIn('\n' + link + '\n', text)  # eigene Zeile, damit WhatsApp ihn als Link erkennt
        self.assertIn('mit Fiete ', self.wa.fotoeinwilligung_nachricht('A', ['Fiete'], link))
        self.assertNotIn(' mit ', self.wa.fotoeinwilligung_nachricht('A', [], link).split('\n')[2])

    def test_nachricht_nennt_alle_punkte(self):
        text = self.wa.fehlende_daten_nachricht('Jürgen', 'Bello', ['Impfung L: Datum fehlt noch', 'Geburtstag'])
        self.assertTrue(text.startswith('Hallo Jürgen,'))
        self.assertIn('für Bello', text)
        self.assertIn('• Impfung L: Datum fehlt noch\n• Geburtstag', text)

    def test_kontaktdaten_nachricht_an_ansprechpartner(self):
        text = self.wa.kontaktdaten_nachricht('Erika Beispiel', 'Jürgen Müller', ['Bello', 'Luna'])
        self.assertTrue(text.startswith('Hallo Erika,'))
        self.assertIn('von Jürgen Müller (mit Bello und Luna) habe ich im Hundemanager weder eine '
                      'Handynummer noch eine E-Mail-Adresse', text)
        self.assertIn('von Anna Neu habe ich', self.wa.kontaktdaten_nachricht('Erika', 'Anna Neu', []))


class WhatsAppOberflaecheTests(unittest.TestCase):
    """Handynummer und E-Mail beim Halter, WhatsApp-Knöpfe in der Übersicht."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def person_speichern(self, person_id, **felder):
        werte = {'vorname': 'Jürgen', 'nachname': 'Müller'}
        werte.update(felder)
        return self.oeffne(f'/person/{person_id}/bearbeiten', urllib.parse.urlencode(werte).encode()).decode('utf-8')

    def test_handynummer_und_email_am_halter(self):
        # Die Datenbank ist von vor 5.1.4 - die neuen Spalten ergänzt die Migration
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('Handynummer fehlt', seite)
        self.assertNotIn('web.whatsapp.com', seite)
        self.assertIn('Für eine WhatsApp-Nachfrage fehlt die Handynummer', seite)

        seite = self.person_speichern(1, mobil='0171 12a', email='juergen@example.org')
        self.assertIn('Die Handynummer stimmt so nicht', seite)
        self.assertIn('value="0171 12a"', seite)  # Eingabe bleibt stehen
        seite = self.person_speichern(1, mobil='0171 1234567', email='kein-at-zeichen')
        self.assertIn('Die E-Mail-Adresse stimmt so nicht', seite)

        seite = self.person_speichern(1, mobil='0171 1234567', email='juergen@example.org', fotofreigabe='1')
        self.assertIn('Person wurde aktualisiert', seite)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT mobil, email, fotofreigabe FROM person WHERE id = 1').fetchone(),
                             ('0171 1234567', 'juergen@example.org', 1))
        self.assertIn('value="0171 1234567"', self.oeffne('/person/1/bearbeiten').decode('utf-8'))

        # Neue Person mit Nummer
        neu = urllib.parse.urlencode({'vorname': 'Anna', 'nachname': 'Neu', 'mobil': '+43 664 1234567'}).encode()
        self.assertIn('Person &#34;Anna Neu&#34; wurde angelegt', self.oeffne('/person/neu', neu).decode('utf-8'))

        import openpyxl
        mappe = openpyxl.load_workbook(io.BytesIO(self.oeffne('/export/excel')))
        zeilen = list(mappe.active.iter_rows(values_only=True))
        self.assertEqual(zeilen[0][11:13], ('Handynummer', 'E-Mail'))
        self.assertIn(('0171 1234567', 'juergen@example.org'), [z[11:13] for z in zeilen[1:]])

    def test_rufname_speichern_und_in_nachrichten_verwenden(self):
        self.person_speichern(1, rufname='  Jürgi  ', mobil='0171 1234567', email='juergen@example.org')
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT vorname, rufname FROM person WHERE id = 1').fetchone(),
                             ('Jürgen', 'Jürgi'))
        self.assertIn('value="Jürgi"', self.oeffne('/person/1/bearbeiten').decode('utf-8'))

        def nachrichtentexte():
            seiten = [self.oeffne('/').decode('utf-8'),
                      self.oeffne('/hund/2/nachweise?art=impfpass').decode('utf-8')]
            texte = []
            for seite in seiten:
                for link in re.findall(r'href="([^"]+)"', seite):
                    link = html_unescape(link)
                    if link.startswith(('https://web.whatsapp.com/', 'mailto:')):
                        abfrage = urllib.parse.parse_qs(urllib.parse.urlparse(link).query)
                        texte.extend(abfrage.get('text', []) + abfrage.get('body', []))
            return texte

        texte = nachrichtentexte()
        for kanal in ('WhatsApp', 'E-Mail'):
            for betreff in ('fehlen mir noch', 'Impfpass etwas', 'Einwilligung'):
                self.assertTrue(any(betreff in t and kanal in t for t in texte), (kanal, betreff))
        self.assertTrue(all(t.startswith('Hallo Jürgi,') for t in texte))

        self.person_speichern(1, rufname='   ', mobil='0171 1234567', email='juergen@example.org')
        with closing(sqlite3.connect(self.db)) as v:
            self.assertIsNone(v.execute('SELECT rufname FROM person WHERE id = 1').fetchone()[0])
        self.assertTrue(all(t.startswith('Hallo Jürgen,') for t in nachrichtentexte()))

        neu = urllib.parse.urlencode({'vorname': 'Franziska', 'nachname': 'Neu', 'rufname': 'Franzi'}).encode()
        self.oeffne('/person/neu', neu)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute("SELECT rufname FROM person WHERE vorname = 'Franziska'").fetchone(),
                             ('Franzi',))
        seite = self.person_speichern(1, rufname='Jürgi', email='ungueltig')
        self.assertIn('value="Jürgi"', seite)

    def test_whatsapp_knoepfe_in_der_uebersicht(self):
        self.person_speichern(1, mobil='0171 1234567')
        seite = self.oeffne('/').decode('utf-8')
        links = [html_unescape(l) for l in re.findall(r'href="(https://web\.whatsapp\.com/[^"]+)"', seite)]

        # Freie Nachricht an den Halter: Chat ohne Text
        self.assertIn('https://web.whatsapp.com/send?phone=491711234567', links)
        # Bello hat SHP/DAP/DHP bis 2027 und Geburtstag, Luna hat gar nichts
        texte = {}
        for l in links:
            abfrage = urllib.parse.parse_qs(urllib.parse.urlparse(l).query)
            if 'text' in abfrage and 'Einwilligung' not in abfrage['text'][0]:
                texte[re.search(r'für (\w+)', abfrage['text'][0]).group(1)] = abfrage['text'][0]
        self.assertEqual(set(texte), {'Bello', 'Luna'})  # Rex/Fiete: Halter ohne Nummer
        self.assertIn('Hallo Jürgen,', texte['Luna'])
        self.assertIn('• Impfung SHP/DAP/DHP: Datum fehlt noch', texte['Luna'])
        self.assertIn('• Haftpflicht: Versicherungsnachweis fehlt noch', texte['Luna'])
        self.assertIn('• Geburtstag', texte['Luna'])
        self.assertNotIn('SHP/DAP/DHP', texte['Bello'])
        self.assertNotIn('Geburtstag', texte['Bello'])
        self.assertIn('• Impfung L: Datum fehlt noch', texte['Bello'])
        self.assertIn('target="whatsapp"', seite)

    def test_kontaktdaten_beim_ansprechpartner_anfragen(self):
        def einstellen(**werte):
            return self.oeffne('/einstellungen', urllib.parse.urlencode(werte).encode()).decode('utf-8')

        # Ohne Handynummer und E-Mail und ohne Ansprechpartner: Knopf führt zu den Einstellungen
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('erst Ansprechpartner eintragen', seite)
        self.assertIn('id="einstellung-kontaktdaten_name"', seite)
        self.assertNotIn('web.whatsapp.com', seite)

        self.assertIn('bitte Name und Handynummer angeben', einstellen(kontaktdaten_name='Erika Beispiel'))
        self.assertIn('Handynummer: stimmt so nicht', einstellen(kontaktdaten_name='Erika Beispiel',
                                                                  kontaktdaten_mobil='12a'))
        self.assertIn('Einstellungen gespeichert', einstellen(kontaktdaten_name='Erika Beispiel',
                                                              kontaktdaten_mobil='0151 2345678'))

        # Der Dialog zeigt, wie die Anfrage ankommt - mit dem Vornamen des Ansprechpartners
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('Hallo <span data-kontakt="vorname">Erika</span>,', seite)
        self.assertIn('<span data-kontakt="mobil">0151 2345678</span>', seite)
        self.assertRegex(seite, r'class="einst-hinweis" data-kontakt="hinweis"\s+hidden')

        seite = self.oeffne('/').decode('utf-8')
        self.assertNotIn('erst Ansprechpartner eintragen', seite)
        self.assertIn('Kontaktdaten bei Erika anfragen', seite)
        texte = [urllib.parse.parse_qs(urllib.parse.urlparse(html_unescape(l)).query)['text'][0]
                 for l in re.findall(r'href="(https://web\.whatsapp\.com/send\?phone=491512345678[^"]+)"', seite)]
        self.assertIn('Hallo Erika,', texte[0])
        self.assertTrue(any('von Jürgen Müller (mit Bello und Luna) habe ich' in t for t in texte), texte)
        self.assertIn('Kontaktdaten bei Erika anfragen', self.oeffne('/person/1/fotoeinwilligung').decode('utf-8'))

        # Mit E-Mail ist der Halter erreichbar - keine Anfrage mehr für ihn
        self.person_speichern(1, email='juergen@example.org')
        seite = self.oeffne('/').decode('utf-8')
        self.assertNotIn('von Jürgen Müller', html_unescape(urllib.parse.unquote(seite)))


class PassnummerTests(unittest.TestCase):
    """Nummer des EU-Heimtierausweises am Hund - erfundene Nummern, nie echte."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def hund_speichern(self, hund_id, **felder):
        werte = {'name': 'Bello', 'geburtstag': '2020-01-01'}
        werte.update(felder)
        return self.oeffne(f'/hund/{hund_id}/bearbeiten', urllib.parse.urlencode(werte).encode()).decode('utf-8')

    def test_passnummer_speichern_anzeigen_exportieren(self):
        # Die Test-Datenbank hat noch keine Spalte passnummer - die Migration ergänzt sie
        self.assertIn('Hund &#34;Bello&#34; wurde aktualisiert', self.hund_speichern(1, passnummer=' de12 345-6789 '))
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT passnummer FROM hund WHERE id = 1').fetchone(), ('DE123456789',))
        self.assertIn('value="DE12 3456789"', self.oeffne('/hund/1/bearbeiten').decode('utf-8'))
        self.assertIn('Pass DE12 3456789', self.oeffne('/').decode('utf-8'))

        import openpyxl
        zeilen = list(openpyxl.load_workbook(io.BytesIO(self.oeffne('/export/excel'))).active.iter_rows(values_only=True))
        self.assertEqual(zeilen[0][13], 'Passnummer')
        self.assertIn(('Bello', 'DE12 3456789'), [(z[3], z[13]) for z in zeilen[1:]])

        # Leeres Feld entfernt die Nummer wieder
        self.hund_speichern(1, passnummer='  ')
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT passnummer FROM hund WHERE id = 1').fetchone(), (None,))


class MailtoTests(unittest.TestCase):
    """mailto:-Links: Empfänger, Betreff und Text so codiert, dass GMX MailCheck sie übernimmt."""

    def setUp(self):
        self.mt = lade_modul(os.path.join(REPO, 'mailto.py'), 'mailto_test')

    def test_link_ohne_und_mit_text(self):
        self.assertEqual(self.mt.link(' juergen@example.org '), 'mailto:juergen@example.org')
        for falsch in ('', None, 'kein-at-zeichen', 'a b@example.org', 'a@b'):
            self.assertIsNone(self.mt.link(falsch), falsch)
        text = 'Hallo Zoë,\nImpfung SHP/DAP/DHP & L: 50 %? a+b'
        link = self.mt.link('juergen@example.org', 'Bello: Impfpass', text)
        self.assertTrue(link.startswith('mailto:juergen@example.org?'))
        self.assertNotIn(' ', link)
        self.assertNotIn('+', link)  # sonst liest der Browser ein Leerzeichen
        # So liest MailCheck die Parameter (URLSearchParams)
        abfrage = urllib.parse.parse_qs(urllib.parse.urlsplit(link).query)
        self.assertEqual(abfrage['subject'], ['Bello: Impfpass'])
        self.assertEqual(abfrage['body'], [text.replace('\n', '\r\n')])

    def test_texte_fuer_email(self):
        wa = lade_modul(os.path.join(REPO, 'whatsapp.py'), 'whatsapp_test')
        for text in (wa.fehlende_daten_nachricht('A', 'Bello', ['Geburtstag'], 'email'),
                     wa.impfpass_nachricht('A', 'Bello', ['Impfung L: Datum fehlt noch'], 'email'),
                     wa.fotoeinwilligung_nachricht('A', ['Bello'], 'https://example.org/f.pdf', 'email')):
            self.assertIn('als Antwort auf diese E-Mail', text)
            self.assertNotIn('WhatsApp', text)
        self.assertIn('hier per WhatsApp', wa.fehlende_daten_nachricht('A', 'Bello', ['Geburtstag']))


class EmailOberflaecheTests(unittest.TestCase):
    """E-Mail-Knöpfe neben WhatsApp und die GMX-Hilfe, die nach ein paar Klicks verschwindet."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def mailto_links(self, seite):
        return [html_unescape(l) for l in re.findall(r'href="(mailto:[^"]+)"', seite)]

    def test_email_knoepfe_und_hilfe(self):
        seite = self.oeffne('/').decode('utf-8')
        self.assertEqual(self.mailto_links(seite), [])
        self.assertNotIn('? GMX-Hilfe', seite)  # ohne E-Mail-Knopf keine Hilfe daneben
        self.assertIn('Hilfe: E-Mail', seite)  # die Anleitung bleibt unten erreichbar
        self.assertIn('Für eine WhatsApp-Nachfrage fehlt die Handynummer', seite)

        # Nur E-Mail, keine Handynummer
        self.oeffne('/person/1/bearbeiten', urllib.parse.urlencode(
            {'vorname': 'Jürgen', 'nachname': 'Müller', 'email': 'juergen@example.org'}).encode())
        seite = self.oeffne('/').decode('utf-8')
        links = self.mailto_links(seite)
        self.assertIn('mailto:juergen@example.org', links)  # freie Nachricht
        texte = {}
        for l in links:
            abfrage = urllib.parse.parse_qs(urllib.parse.urlsplit(l).query)
            if 'subject' in abfrage:
                texte[abfrage['subject'][0]] = abfrage['body'][0]
        self.assertEqual(set(texte), {'Bello: fehlende Angaben', 'Luna: fehlende Angaben',
                                      'Einwilligung zu Fotoaufnahmen'})
        self.assertIn('Hallo Jürgen,', texte['Luna: fehlende Angaben'])
        self.assertIn('• Geburtstag', texte['Luna: fehlende Angaben'])
        self.assertIn('als Antwort auf diese E-Mail', texte['Luna: fehlende Angaben'])
        self.assertIn('per E-Mail anfragen', seite)
        self.assertIn('? GMX-Hilfe', seite)
        self.assertIn('E-Mail Links in Webseiten mit MailCheck öffnen', seite)

        # Impfpass-Seite und fällige Impfungen
        seite = self.oeffne('/hund/2/nachweise?art=impfpass').decode('utf-8')
        self.assertIn('Foto vom Impfpass per E-Mail', seite)
        self.assertEqual([urllib.parse.parse_qs(urllib.parse.urlsplit(l).query)['subject'][0]
                          for l in self.mailto_links(seite)], ['Luna: Impfpass'])
        self.assertNotIn('fehlt die Handynummer', seite)

        # Nach drei geklickten E-Mail-Links verschwindet die Hilfe neben den Knöpfen
        for _ in range(3):
            self.assertIn('? GMX-Hilfe', self.oeffne('/').decode('utf-8'))
            self.oeffne('/email/geklickt', b'')
        seite = self.oeffne('/').decode('utf-8')
        self.assertNotIn('? GMX-Hilfe', seite)
        self.assertIn('Hilfe: E-Mail', seite)
        self.oeffne('/email/geklickt', b'')
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute("SELECT wert FROM einstellung WHERE schluessel = 'mail_klicks'").fetchone(),
                             ('3',))
        # Der Zähler gehört nicht in den Einstellungsdialog und übersteht "Standard wiederherstellen"
        self.assertNotIn('mail_klicks', self.oeffne('/einstellungen').decode('utf-8'))
        self.oeffne('/einstellungen', urllib.parse.urlencode({'standard': '1'}).encode())
        self.assertNotIn('? GMX-Hilfe', self.oeffne('/').decode('utf-8'))


class WiedervorlageTests(unittest.TestCase):
    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def datum(self):
        with closing(sqlite3.connect(self.db)) as v:
            return v.execute('SELECT wiedervorlage_am FROM hund WHERE id = 1').fetchone()[0]

    def test_eigenes_datum_speichern_anzeigen_und_loeschen(self):
        termin = (date.today() + timedelta(days=40)).isoformat()
        ergebnis = json.loads(self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(
            dict(aktion='setzen', datum=termin)).encode()))
        self.assertEqual(self.datum(), termin)
        self.assertEqual(ergebnis['datum_iso'], termin)
        self.assertFalse(ergebnis['faellig'])
        for pfad in ('/', '/hund/1/nachweise?art=impfpass'):
            seite = self.oeffne(pfad).decode('utf-8')
            self.assertIn(date.fromisoformat(termin).strftime('%d.%m.%Y'), seite)
            self.assertIn(f'value="{termin}"', seite)
            self.assertIn('data-wiedervorlage-loeschen', seite)
        for _ in range(2):
            ergebnis = json.loads(self.oeffne('/hund/1/wiedervorlage', b'aktion=loeschen'))
            self.assertIsNone(self.datum())
            self.assertIsNone(ergebnis['datum'])
            self.assertIsNone(ergebnis['datum_iso'])
            self.assertFalse(ergebnis['faellig'])
        self.assertNotIn(f'value="{termin}"', self.oeffne('/').decode('utf-8'))
        self.oeffne('/hund/1/wiedervorlage', b'')
        self.assertEqual(self.datum(), (date.today() + timedelta(days=14)).isoformat())

    def test_ungueltige_aktionen_lassen_termin_unveraendert(self):
        self.oeffne('/hund/1/wiedervorlage', b'')
        vorher = self.datum()
        for daten in ({'aktion': 'setzen'}, {'aktion': 'anders'}, {'aktion': ''},
                      *({'aktion': 'setzen', 'datum': wert} for wert in
                        ('', 'abc', '2026-02-30', '08.10.2026', '20261008', '2026-1-1'))):
            with self.subTest(daten=daten):
                with self.assertRaises(urllib.error.HTTPError) as fehler:
                    self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(daten).encode())
                self.assertEqual(fehler.exception.code, 400)
                self.assertTrue(json.loads(fehler.exception.read())['fehler'])
                self.assertEqual(self.datum(), vorher)
        with self.assertRaises(urllib.error.HTTPError) as fehler:
            self.oeffne('/hund/99999/wiedervorlage', b'aktion=loeschen')
        self.assertEqual(fehler.exception.code, 404)

    def test_zukuenftiger_termin_stellt_nur_handlungsbedarf_zurueck(self):
        gestern = (date.today() - timedelta(days=1)).isoformat()
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE hund SET gueltig_l = ? WHERE id = 1', (gestern,))
            v.commit()
        morgen = (date.today() + timedelta(days=1)).isoformat()
        for daten in (b'', urllib.parse.urlencode(dict(aktion='setzen', datum=morgen)).encode()):
            self.oeffne('/hund/1/wiedervorlage', daten)
            seite = self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8')
            self.assertFalse('data-wiedervorlage="1"' in seite)
            self.assertIn('Handlungsbedarf (3)', seite)
            for pfad in ('/', '/?ansicht=alle', '/?filter=abgelaufen', '/nachweise?art=impfpass'):
                self.assertIn('data-wiedervorlage="1"', self.oeffne(pfad).decode('utf-8'))
        for termin in (date.today().isoformat(), gestern):
            ergebnis = json.loads(self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(
                dict(aktion='setzen', datum=termin)).encode()))
            self.assertTrue(ergebnis['faellig'])
            seite = self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8')
            self.assertIn('data-wiedervorlage="1"', seite)
            self.assertIn('Handlungsbedarf (4)', seite)
        self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(
            dict(aktion='setzen', datum=morgen)).encode())
        self.oeffne('/hund/1/wiedervorlage', b'aktion=loeschen')
        self.assertIn('Handlungsbedarf (4)', self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8'))
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE person SET aktiv = 0 WHERE id = 1')
            v.commit()
        seite = self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8')
        self.assertNotIn('data-wiedervorlage="1"', seite)
        self.assertIn('Handlungsbedarf (2)', seite)

    def test_speichern_einstellungen_und_faellige_anzeige(self):
        self.assertIsNone(self.datum())  # Migration aus altem Schema, keine erfundene Erinnerung
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('data-wartezeit="30"', seite)
        ergebnis = json.loads(self.oeffne('/hund/1/wiedervorlage', b''))
        termin = date.today() + timedelta(days=14)
        self.assertEqual(self.datum(), termin.isoformat())
        self.assertEqual(ergebnis['datum'], termin.strftime('%d.%m.%Y'))
        self.assertIn('Wiedervorlage am', self.oeffne('/').decode('utf-8'))
        self.oeffne('/einstellungen', urllib.parse.urlencode(
            {'wiedervorlage_tage': '7', 'wiedervorlage_wartezeit': '45'}).encode())
        self.oeffne('/hund/1/wiedervorlage', b'')
        self.assertEqual(self.datum(), (date.today() + timedelta(days=7)).isoformat())
        self.assertIn('data-wartezeit="45"', self.oeffne('/').decode('utf-8'))
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE hund SET wiedervorlage_am = ? WHERE id = 1', (date.today().isoformat(),))
            v.commit()
        self.assertIn('Wiedervorlage fällig seit', self.oeffne('/').decode('utf-8'))
        self.oeffne('/einstellungen', b'standard=1')
        self.assertIn('data-wartezeit="30"', self.oeffne('/').decode('utf-8'))

    def test_ungueltige_einstellungen_werden_abgelehnt(self):
        for wert in ('0', '-1', '1.5', 'abc', '999999'):
            seite = self.oeffne('/einstellungen', urllib.parse.urlencode(
                {'wiedervorlage_tage': wert, 'wiedervorlage_wartezeit': wert}).encode()).decode('utf-8')
            self.assertIn('bitte eine ganze Zahl', seite)
        self.oeffne('/hund/1/wiedervorlage', b'')
        self.assertEqual(self.datum(), (date.today() + timedelta(days=14)).isoformat())

    def test_nur_vollstaendige_angaben_entfernen_wiedervorlage(self):
        self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(
            dict(aktion='setzen', datum=(date.today() + timedelta(days=40)).isoformat())).encode())
        zukunft = (date.today() + timedelta(days=365)).isoformat()
        daten = dict(name='Bello', geburtstag='', haftpflicht_gueltig='ja',
                     gueltig_shp_dap_dhp=zukunft, gueltig_l=zukunft,
                     gueltig_bbpi=zukunft, gueltig_t=zukunft)
        self.oeffne('/hund/1/bearbeiten', urllib.parse.urlencode(daten).encode())
        self.assertIsNotNone(self.datum())  # Häkchen ersetzt keinen Versicherungsnachweis
        koerper, typ = multipart('datei', 'Test.pdf', test_pdf(bestaetigung_zeilen(date.today())))
        seite = self.oeffne('/hund/1/haftpflicht/hochladen', koerper, typ).decode('utf-8')
        sha = re.search(r'name="sha256" value="([0-9a-f]{64})"', seite).group(1)
        self.oeffne('/hund/1/haftpflicht/speichern', urllib.parse.urlencode(
            dict(sha256=sha, endung='pdf', original_name='Test.pdf', gueltig_bis=zukunft)).encode())
        self.assertIsNotNone(self.datum())  # Geburtstag fehlt noch
        daten['geburtstag'] = '2020-01-01'
        self.oeffne('/hund/1/bearbeiten', urllib.parse.urlencode(daten).encode())
        self.assertIsNone(self.datum())
        self.assertIsNone(json.loads(self.oeffne('/hund/1/wiedervorlage', b''))['datum'])
        self.assertIsNone(json.loads(self.oeffne('/hund/1/wiedervorlage', urllib.parse.urlencode(
            dict(aktion='setzen', datum=zukunft)).encode()))['datum'])

    def test_faellige_wiedervorlage_bleibt_im_handlungsbedarf(self):
        zukunft = (date.today() + timedelta(days=365)).isoformat()
        with closing(sqlite3.connect(self.db)) as v:
            v.execute('UPDATE hund SET geburtstag = NULL, gueltig_shp_dap_dhp = ?, gueltig_l = ?, '
                      'gueltig_bbpi = ?, gueltig_t = ?, wiedervorlage_am = ? WHERE id = 1',
                      (zukunft, zukunft, zukunft, zukunft, date.today().isoformat()))
            v.execute("INSERT INTO nachweis (hund_id, art, datei_sha256, datei_endung, original_name, "
                      "hochgeladen_am, gueltig_bis) VALUES (1, 'haftpflicht', ?, 'pdf', 'Test.pdf', ?, ?)",
                      ('a' * 64, date.today().isoformat(), zukunft))
            v.commit()
        seite = self.oeffne('/?ansicht=handlungsbedarf').decode('utf-8')
        self.assertIn('Wiedervorlage fällig seit', seite)
        self.assertIn('data-wiedervorlage="1"', seite)

    def test_anfragen_sind_mit_countdown_verbunden(self):
        self.oeffne('/person/1/bearbeiten', urllib.parse.urlencode(
            dict(vorname='Jürgen', nachname='Müller', mobil='0171 1234567', email='test@example.org')).encode())
        for pfad in ('/', '/hund/1/nachweise?art=impfpass'):
            seite = self.oeffne(pfad).decode('utf-8')
            self.assertIn('data-wiedervorlage-link="1"', seite)
            self.assertIn('data-wiedervorlage="1"', seite)
            self.assertIn('data-wiedervorlage-abbrechen', seite)
            self.assertIn('data-wiedervorlage-ok', seite)
        self.assertIsNone(self.datum())  # Öffnen allein speichert keine Wiedervorlage


class FotoeinwilligungTests(unittest.TestCase):
    """Fotoeinwilligung je Halter: Anfrage per WhatsApp mit Formular-Link, Nachweis, Widerruf."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne
    STANDARD_LINK = 'https://www.hsv-grossbottwar.de/wp-content/uploads/2026/02/Einwilligung_Fotoaufnahmen_wolf.pdf'

    def anfragen(self):
        """Texte der Fotoeinwilligungs-Anfragen in der Übersicht."""
        seite = self.oeffne('/').decode('utf-8')
        texte = []
        for l in re.findall(r'href="(https://web\.whatsapp\.com/[^"]+)"', seite):
            abfrage = urllib.parse.parse_qs(urllib.parse.urlparse(html_unescape(l)).query)
            if 'Einwilligung' in abfrage.get('text', [''])[0]:
                texte.append(abfrage['text'][0])
        return texte

    def einstellen(self, **werte):
        return self.oeffne('/einstellungen', urllib.parse.urlencode(werte).encode()).decode('utf-8')

    def hochladen(self, person_id, daten, dateiname, **felder):
        koerper, typ = multipart('datei', dateiname, daten)
        grenze = typ.split('boundary=')[1]
        for name, wert in felder.items():
            koerper = (f'--{grenze}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{wert}\r\n').encode() + koerper
        return self.oeffne(f'/person/{person_id}/fotoeinwilligung/hochladen', koerper, typ).decode('utf-8')

    def excel_foto(self):
        import openpyxl
        mappe = openpyxl.load_workbook(io.BytesIO(self.oeffne('/export/excel')))
        return {z[1]: z[2] for z in mappe.active.iter_rows(min_row=2, values_only=True)}

    def test_anfrage_nachweis_und_widerruf(self):
        # Ohne Einwilligung: Etikett, Kennzahl und - mit Handynummer - die Anfrage mit Formular-Link
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('Keine Fotoeinwilligung', seite)
        self.assertIn('<strong>3</strong>Halter ohne unterschriebene Fotoeinwilligung', seite)
        self.assertEqual(self.anfragen(), [])  # noch keine Handynummern
        self.oeffne('/person/1/bearbeiten', urllib.parse.urlencode(
            {'vorname': 'Jürgen', 'nachname': 'Müller', 'mobil': '0171 1234567'}).encode())
        [text] = self.anfragen()
        self.assertIn('Hallo Jürgen,', text)
        self.assertIn('mit Bello und Luna', text)
        self.assertIn(self.STANDARD_LINK, text)
        self.assertEqual(self.excel_foto()['Jürgen'], 'Nein')

        # Unterschriebene Einwilligung als Foto hochladen
        foto = b'\xff\xd8\xff\xe0' + b'Einwilligung-Test' * 40
        seite = self.hochladen(1, foto, 'Einwilligung Müller.jpg', unterschrieben_am='2026-02-14',
                               bemerkung='Test')
        self.assertIn('Fotoeinwilligung von Jürgen Müller gespeichert', seite)
        self.assertIn('unterschrieben am 14.02.2026', seite)
        self.assertEqual(self.oeffne('/fotoeinwilligung/1/datei'), foto)
        self.assertEqual(self.anfragen(), [])  # nichts mehr zu fragen
        self.assertIn('Fotoeinwilligung ✓', self.oeffne('/').decode('utf-8'))
        self.assertEqual(self.excel_foto()['Jürgen'], 'Ja')
        # Datei liegt in der gemeinsamen, gesicherten Nachweis-Ablage
        sha = __import__('hashlib').sha256(foto).hexdigest()
        self.assertTrue(os.path.exists(os.path.join(self.app_dir, 'instance', 'nachweise', sha + '.jpg')))

        # Person speichern ohne Häkchen nimmt die belegte Freigabe nicht weg
        seite = self.oeffne('/person/1/bearbeiten', urllib.parse.urlencode(
            {'vorname': 'Jürgen', 'nachname': 'Müller', 'mobil': '0171 1234567'}).encode()).decode('utf-8')
        self.assertIn('Person wurde aktualisiert', seite)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT fotofreigabe FROM person WHERE id = 1').fetchone(), (1,))

        # Widerruf: vermerkt, Dokument bleibt, keine erneute Anfrage
        seite = self.oeffne('/fotoeinwilligung/1/widerrufen', urllib.parse.urlencode(
            {'widerrufen_am': '2026-09-01'}).encode()).decode('utf-8')
        self.assertIn('Widerruf vom 01.09.2026 vermerkt', seite)
        self.assertEqual(self.oeffne('/fotoeinwilligung/1/datei'), foto)
        self.assertIn('Fotoeinwilligung widerrufen', self.oeffne('/').decode('utf-8'))
        self.assertEqual(self.anfragen(), [])
        self.assertEqual(self.excel_foto()['Jürgen'], 'Widerrufen')
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT fotofreigabe FROM person WHERE id = 1').fetchone(), (0,))

    def test_falsche_datei_und_datum_in_der_zukunft(self):
        seite = self.hochladen(2, b'MZ\x90\x00', 'virus.exe')
        self.assertIn('Bitte ein PDF oder ein Foto', seite)
        morgen = (date.today() + timedelta(days=1)).isoformat()
        seite = self.hochladen(2, test_pdf(['Einwilligung']), 'e.pdf', unterschrieben_am=morgen)
        self.assertIn('liegt in der Zukunft', seite)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT COUNT(*) FROM fotoeinwilligung').fetchone(), (0,))

    def test_formular_link_im_einstellungsdialog(self):
        self.oeffne('/person/2/bearbeiten', urllib.parse.urlencode(
            {'vorname': 'Saskia', 'nachname': 'Test', 'mobil': '0171 7654321'}).encode())
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('id="einstellungen-dialog"', seite)  # Dialog auf jeder Seite
        self.assertIn(f'value="{self.STANDARD_LINK}"', seite)
        # Dokument-Karte: Dateiname statt langer Adresse, Eingabefeld erst auf Klick
        self.assertIn('Einwilligung_Fotoaufnahmen_wolf.pdf</strong>', seite)
        self.assertIn('<details class="link-aendern">', seite)

        # Unsinn wird abgelehnt, nichts gespeichert
        seite = self.einstellen(fotoeinwilligung_link='kein link')
        self.assertIn('bitte eine vollständige Adresse', seite)
        self.assertIn(self.STANDARD_LINK, self.anfragen()[0])

        # Neues Formular: Link ändern -> Anfrage nutzt ihn, Rücksprung auf die Seite davor
        neu = 'https://www.example.org/formulare/Einwilligung_2027.pdf'
        seite = self.einstellen(fotoeinwilligung_link=neu, zurueck='/?ansicht=alle')
        self.assertIn('Einstellungen gespeichert', seite)
        self.assertIn('Alle Hunde', seite)
        [text] = self.anfragen()
        self.assertIn(neu, text)
        self.assertNotIn(self.STANDARD_LINK, text)
        self.assertIn(neu, self.oeffne('/person/2/fotoeinwilligung').decode('utf-8'))
        # Fremde Rücksprung-Ziele werden ignoriert
        self.assertIn('Einstellungen gespeichert', self.einstellen(fotoeinwilligung_link=neu, zurueck='//boese.example'))

        # Zurück auf den Standard
        self.einstellen(standard='1', fotoeinwilligung_link=neu)
        self.assertIn(self.STANDARD_LINK, self.anfragen()[0])
        self.assertIn(self.STANDARD_LINK, self.oeffne('/einstellungen').decode('utf-8'))


class ThemeTests(unittest.TestCase):
    """Farbschema in den Einstellungen: Türkis als Standard, das bisherige Grün wählbar."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne
    einstellen = FotoeinwilligungTests.einstellen

    def gespeichert(self):
        with closing(sqlite3.connect(self.db)) as v:
            zeile = v.execute("SELECT wert FROM einstellung WHERE schluessel = 'theme'").fetchone()
        return zeile and zeile[0]

    def test_theme_waehlen(self):
        seite = self.oeffne('/').decode('utf-8')
        self.assertIn('data-theme="tuerkis"', seite)
        self.assertRegex(seite, r'<input type="radio" name="theme" value="tuerkis" id="einstellung-theme-tuerkis"\s+checked>')
        self.assertIsNone(self.gespeichert())  # Standard wird nicht gespeichert

        self.assertIn('Einstellungen gespeichert', self.einstellen(theme='gruen'))
        for pfad in ('/', '/nachweise', '/einstellungen', '/sicherungen'):
            self.assertIn('data-theme="gruen"', self.oeffne(pfad).decode('utf-8'))
        self.assertEqual(self.gespeichert(), 'gruen')

        # Unbekanntes Schema wird abgelehnt, das gewählte bleibt
        self.assertIn('Farbschema: bitte', self.einstellen(theme='pink'))
        self.assertEqual(self.gespeichert(), 'gruen')

        self.einstellen(standard='1', theme='gruen')
        self.assertIn('data-theme="tuerkis"', self.oeffne('/').decode('utf-8'))
        self.assertIsNone(self.gespeichert())


class BilderTests(unittest.TestCase):
    """Fotos richtig herum - erkannt am Inhalt, gedreht wird nur die Anzeige."""

    @classmethod
    def setUpClass(cls):
        cls.bilder = lade_modul(os.path.join(REPO, 'bilder.py'), 'bilder_test')

    def test_quer_und_kopfueber_fotografierte_seiten_werden_erkannt(self):
        for art in ('block', 'felder'):
            for verdreht in (0, 90, 180, 270):
                with self.subTest(art=art, verdreht=verdreht):
                    noetig = (360 - verdreht) % 360
                    self.assertEqual(self.bilder.ausrichtung_erkennen(textfoto(art, verdreht)), noetig)
                    self.assertEqual(self.bilder.ausrichtung_erkennen(textfoto(art, verdreht, 'PNG')), noetig)

    def test_ohne_text_oder_ohne_bild_wird_nicht_gedreht(self):
        from PIL import Image
        puffer = io.BytesIO()
        Image.linear_gradient('L').resize((800, 600)).convert('RGB').save(puffer, 'JPEG')
        self.assertEqual(self.bilder.ausrichtung_erkennen(puffer.getvalue()), 0)
        self.assertEqual(self.bilder.ausrichtung_erkennen(test_pdf(['Kein Foto'])), 0)
        self.assertEqual(self.bilder.ausrichtung_erkennen(b'\xff\xd8\xff kaputt'), 0)
        self.assertIsNone(self.bilder.ansicht(test_pdf(['Kein Foto'])))

    def test_ansicht_dreht_und_beachtet_kamera_ausrichtung(self):
        from PIL import Image
        foto = textfoto()  # 1400 x 1000
        daten, typ = self.bilder.ansicht(foto, 90)
        self.assertEqual(typ, 'image/jpeg')
        self.assertEqual(Image.open(io.BytesIO(daten)).size, (1000, 1400))
        # Handy speichert "um 90 Grad drehen" nur als EXIF-Angabe - die Anzeige wendet sie an
        daten, _ = self.bilder.ansicht(textfoto(exif_ausrichtung=6), 0)
        self.assertEqual(Image.open(io.BytesIO(daten)).size, (1000, 1400))
        # PNG mit Transparenz wird zu JPEG auf weißem Grund
        puffer = io.BytesIO()
        Image.new('RGBA', (50, 40), (0, 0, 0, 0)).save(puffer, 'PNG')
        daten, _ = self.bilder.ansicht(puffer.getvalue(), 270)
        self.assertEqual(Image.open(io.BytesIO(daten)).size, (40, 50))

    def test_drehung_aus_formular(self):
        for wert, erwartet in (('90', 90), ('-90', 270), (450, 90), ('45', 0), ('x', 0), (None, 0)):
            self.assertEqual(self.bilder.drehung_pruefen(wert), erwartet)


class ImpfpassOberflaecheTests(unittest.TestCase):
    """Impfpass-Fotos hochladen, Fälligkeiten eintragen, WhatsApp-Nachfrage bei Ablauf."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def hund(self, *spalten):
        with closing(sqlite3.connect(self.db)) as v:
            return v.execute(f'SELECT {", ".join(spalten)} FROM hund WHERE id = 1').fetchone()

    def test_fotos_hochladen_impfungen_eintragen_und_nachfragen(self):
        from PIL import Image
        heute = date.today()
        self.oeffne('/person/1/bearbeiten', urllib.parse.urlencode(
            {'vorname': 'Jürgen', 'nachname': 'Müller', 'mobil': '0171 1234567'}).encode())

        # Zwei Seiten auf einmal - die erste quer fotografiert
        koerper, typ = multipart_mehrere([('art', 'impfpass')], [
            ('datei', 'seite1.jpg', textfoto('block', 90)), ('datei', 'seite2.jpg', textfoto('felder'))])
        seite = self.oeffne('/hund/1/nachweise/hochladen', koerper, typ).decode('utf-8')
        self.assertIn('Seite 1 von 2', seite)
        self.assertIn('automatisch gedreht', seite)
        self.assertEqual(re.findall(r'name="drehung" value="(\d+)"', seite), ['270', '0'])
        shas = re.findall(r'name="sha256" value="([0-9a-f]{64})"', seite)
        self.assertEqual(len(shas), 2)
        self.assertIn('Beim Hund eingetragen:', seite)
        # Die Anzeige kommt gedreht, die Datei bleibt unverändert
        bild = Image.open(io.BytesIO(self.oeffne(f'/nachweis/bild/{shas[0]}.jpg?drehung=270')))
        self.assertEqual(bild.size, (1400, 1000))
        self.assertEqual(self.oeffne(f'/nachweis/vorschau/{shas[0]}.jpg'), textfoto('block', 90))

        l_bis = heute + timedelta(days=300)
        t_bis = heute - timedelta(days=10)  # Tollwut abgelaufen
        formular = [('art', 'impfpass'), ('aussteller', 'Tierarztpraxis Muster'), ('eingereicht_von', 'Jürgen Müller')]
        for sha, name, drehung in zip(shas, ('seite1.jpg', 'seite2.jpg'), ('270', '0')):
            formular += [('sha256', sha), ('endung', 'jpg'), ('original_name', name), ('drehung', drehung)]
        formular += [('impf_gueltig_l', l_bis.isoformat()), ('impf_gueltig_t', t_bis.isoformat()),
                     ('impf_gueltig_shp_dap_dhp', '2026-06-01')]  # Bello hat schon bis 01.01.2027
        seite = self.oeffne('/hund/1/nachweise/speichern', urllib.parse.urlencode(formular).encode()).decode('utf-8')
        self.assertIn('Impfpass-Nachweis für Bello gespeichert', seite)
        self.assertIn(f'L bis {l_bis.strftime("%d.%m.%Y")}', seite)
        self.assertIn(f'T bis {t_bis.strftime("%d.%m.%Y")}', seite)
        self.assertIn('SHP/DAP/DHP (eingetragen ist schon 01.01.2027)', seite)
        self.assertEqual(self.hund('gueltig_shp_dap_dhp', 'gueltig_l', 'gueltig_t'),
                         ('2027-01-01', l_bis.isoformat(), t_bis.isoformat()))
        with closing(sqlite3.connect(self.db)) as v:
            # Der Nachweis selbst gilt bis zur frühesten Impfung darauf
            self.assertEqual(v.execute('SELECT drehung, impfungen, gueltig_bis FROM nachweis').fetchone(),
                             (270, 'SHP/DAP/DHP, L, T', '2026-06-01'))
            self.assertEqual(v.execute('SELECT nr, drehung, original_name FROM nachweis_seite').fetchall(),
                             [(2, 0, 'seite2.jpg')])

        # WhatsApp-Nachfrage auf der Impfpass-Seite des Hundes ...
        seite = self.oeffne('/hund/1/nachweise?art=impfpass').decode('utf-8')
        links = [html_unescape(l) for l in re.findall(r'href="(https://web\.whatsapp\.com/[^"]+)"', seite)]
        self.assertEqual(len(links), 1)
        text = urllib.parse.parse_qs(urllib.parse.urlparse(links[0]).query)['text'][0]
        self.assertIn(f'• Impfung T: abgelaufen am {t_bis.strftime("%d.%m.%Y")}', text)
        self.assertIn('• Impfung BbPi: Datum fehlt noch', text)
        self.assertNotIn('Impfung L:', text)
        self.assertIn('Foto der Impfpass-Seite', text)
        # ... und in der Liste der fälligen Impfungen
        seite = self.oeffne('/nachweise').decode('utf-8')
        self.assertIn('Impfungen abgelaufen oder bald fällig', seite)
        self.assertIn('seit 10 Tagen abgelaufen', seite)
        self.assertIn('WhatsApp-Nachfrage', seite)
        self.assertNotIn('Impfungen abgelaufen', self.oeffne('/nachweise?art=haftpflicht').decode('utf-8'))

        # Ansehen mit allen Seiten, Seite 2 nachträglich drehen
        seite = self.oeffne('/nachweis/1/ansehen').decode('utf-8')
        self.assertIn('Seite 2 von 2', seite)
        self.oeffne('/nachweis/1/drehen', urllib.parse.urlencode({'seite': '1', 'richtung': 'rechts'}).encode())
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT drehung FROM nachweis_seite').fetchone(), (90,))

        # Korrektur: L war ein Tippfehler - auch ein früheres Datum wird übernommen,
        # weil der Wert beim Hund aus diesem Nachweis stammt
        seite = self.oeffne('/nachweis/1/bearbeiten').decode('utf-8')
        self.assertIn(f'value="{l_bis.isoformat()}"', seite)
        l_korrigiert = heute + timedelta(days=200)
        korrektur = [(k, w) for k, w in formular if k not in ('impf_gueltig_l', 'drehung')]
        korrektur += [('impf_gueltig_l', l_korrigiert.isoformat()), ('drehung', '0'), ('drehung', '90')]
        seite = self.oeffne('/nachweis/1/bearbeiten', urllib.parse.urlencode(korrektur).encode()).decode('utf-8')
        self.assertIn('Angaben zum Nachweis aktualisiert', seite)
        self.assertEqual(self.hund('gueltig_l'), (l_korrigiert.isoformat(),))
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT drehung FROM nachweis').fetchone(), (0,))

        # Unmögliches Datum: Fehlermeldung, Eingaben bleiben stehen, nichts wird geändert
        falsch = korrektur + [('impf_gueltig_bbpi', '2099-05-01')]
        seite = self.oeffne('/nachweis/1/bearbeiten', urllib.parse.urlencode(falsch).encode()).decode('utf-8')
        self.assertIn('kann nicht stimmen', seite)
        self.assertIn('value="2099-05-01"', seite)
        self.assertEqual(self.hund('gueltig_bbpi'), (None,))

    def test_aelteres_foto_ueberschreibt_keine_neuere_impfung(self):
        neu = date.today() + timedelta(days=500)
        alt = date.today() + timedelta(days=100)
        for datum, name in ((neu, 'neu.jpg'), (alt, 'alt.jpg')):
            koerper, typ = multipart_mehrere([('art', 'impfpass')], [('datei', name, textfoto('felder', 0) + name.encode())])
            seite = self.oeffne('/hund/1/nachweise/hochladen', koerper, typ).decode('utf-8')
            sha = re.search(r'name="sha256" value="([0-9a-f]{64})"', seite).group(1)
            seite = self.oeffne('/hund/1/nachweise/speichern', urllib.parse.urlencode([
                ('art', 'impfpass'), ('sha256', sha), ('endung', 'jpg'), ('original_name', name),
                ('drehung', '0'), ('impf_gueltig_t', datum.isoformat())]).encode()).decode('utf-8')
        self.assertIn('Nicht übernommen', seite)
        self.assertEqual(self.hund('gueltig_t'), (neu.isoformat(),))
        # Beide Nachweise bleiben in der Historie
        self.assertEqual(self.oeffne('/hund/1/nachweise?art=impfpass').decode('utf-8').count('Korrigieren'), 2)


class ImpfpassUebergangTests(unittest.TestCase):
    """Saskias Datenbank ab 5.1.3 hat schon Nachweise - ohne Drehung, Impfdaten und weitere Seiten."""

    setUp = HaftpflichtOberflaecheTests.setUp
    tearDown = HaftpflichtOberflaecheTests.tearDown
    oeffne = HaftpflichtOberflaecheTests.oeffne

    def vor_dem_start(self):
        self.foto = textfoto('block', 90)
        self.sha = hashlib.sha256(self.foto).hexdigest()
        ablage = os.path.join(self.app_dir, 'instance', 'nachweise')
        os.makedirs(ablage)
        with open(os.path.join(ablage, f'{self.sha}.jpg'), 'wb') as f:
            f.write(self.foto)
        with closing(sqlite3.connect(self.db)) as v:
            v.executescript(f'''
                ALTER TABLE person ADD COLUMN mobil VARCHAR(30);
                ALTER TABLE person ADD COLUMN email VARCHAR(200);
                CREATE TABLE nachweis (id INTEGER PRIMARY KEY, hund_id INTEGER NOT NULL REFERENCES hund(id),
                    art VARCHAR(20) NOT NULL, datei_sha256 VARCHAR(64) NOT NULL, datei_endung VARCHAR(4) NOT NULL,
                    original_name VARCHAR(255), hochgeladen_am DATETIME NOT NULL, eingereicht_von VARCHAR(200),
                    aussteller VARCHAR(200), nummer VARCHAR(100), ausgestellt_am DATE, gueltig_bis DATE,
                    tier_laut_nachweis VARCHAR(100), chipnummer VARCHAR(30), impfungen VARCHAR(100), bemerkung TEXT);
                INSERT INTO nachweis (hund_id, art, datei_sha256, datei_endung, original_name, hochgeladen_am,
                                      aussteller, gueltig_bis, impfungen)
                VALUES (1, 'impfpass', '{self.sha}', 'jpg', 'pass.jpg', '2026-10-01 10:00:00',
                        'Tierarztpraxis Alt', '2027-03-01', 'SHP/DAP/DHP, L');
            ''')
            v.commit()

    def test_alte_nachweise_bleiben_und_lassen_sich_drehen_und_ergaenzen(self):
        with closing(sqlite3.connect(self.db)) as v:
            spalten = {z[1] for z in v.execute('PRAGMA table_info(nachweis)')}
            self.assertTrue({'drehung', 'impf_gueltig'} <= spalten)
            self.assertEqual(v.execute('SELECT drehung, impfungen FROM nachweis').fetchone(), (0, 'SHP/DAP/DHP, L'))
            self.assertEqual(v.execute('SELECT COUNT(*) FROM nachweis_seite').fetchone(), (0,))
        # Vor der Änderung am Schema wurde gesichert
        self.assertTrue(any('vor-migration' in n for n in os.listdir(os.path.join(self.app_dir, 'instance', 'backup'))))

        seite = self.oeffne('/hund/1/nachweise?art=impfpass').decode('utf-8')
        self.assertIn('Tierarztpraxis Alt', seite)
        self.assertIn('SHP/DAP/DHP, L', seite)
        self.assertIn('Dokument', self.oeffne('/nachweis/1/ansehen').decode('utf-8'))
        self.oeffne('/nachweis/1/drehen', urllib.parse.urlencode({'seite': '0', 'richtung': 'links'}).encode())
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT drehung FROM nachweis').fetchone(), (270,))

        # Alter Nachweis nachträglich mit Daten ergänzt - die alten Häkchen bleiben erhalten
        seite = self.oeffne('/nachweis/1/bearbeiten').decode('utf-8')
        self.assertIn('name="impfungen" value="SHP/DAP/DHP"', seite)
        formular = re.findall(r'<input type="hidden" name="(\w+)" value="([^"]*)"', seite)
        formular += [('impf_gueltig_t', '2027-12-09')]
        seite = self.oeffne('/nachweis/1/bearbeiten', urllib.parse.urlencode(formular).encode()).decode('utf-8')
        self.assertIn('T bis 09.12.2027', seite)
        with closing(sqlite3.connect(self.db)) as v:
            self.assertEqual(v.execute('SELECT impfungen, gueltig_bis FROM nachweis').fetchone(),
                             ('SHP/DAP/DHP, L, T', '2027-12-09'))
            self.assertEqual(v.execute('SELECT gueltig_t FROM hund WHERE id = 1').fetchone(), ('2027-12-09',))
        self.assertEqual(self.oeffne('/nachweis/1/datei'), self.foto)  # Datei unverändert


def html_unescape(text):
    import html
    return html.unescape(text)


if __name__ == '__main__':
    unittest.main()
