# app.py - Hauptanwendung
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date
from dateutil.relativedelta import relativedelta
import io
import json
import os
import re
import secrets
import subprocess
import sys
import threading
import time
import urllib.parse
import traceback
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy import func, inspect, text #NEU
import backup
import bilder
import mailto
import nachweise
import updater
import whatsapp

app = Flask(__name__)
app.config['SECRET_KEY'] = secrets.token_hex(32)
# Fester Pfad, damit App und Sicherung garantiert dieselbe Datei meinen
os.makedirs(backup.INSTANCE_DIR, exist_ok=True)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + backup.DB_PFAD
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
# Mehrere Fotos auf einmal (z.B. alle Seiten des Impfpasses) - je Datei gilt MAX_GROESSE
app.config['MAX_CONTENT_LENGTH'] = 4 * nachweise.MAX_GROESSE + 1024 * 1024

db = SQLAlchemy(app)

# Datenbankmodelle
class Person(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    vorname = db.Column(db.String(100), nullable=False)
    nachname = db.Column(db.String(100), nullable=False)
    aktiv = db.Column(db.Boolean, default=True)  # 👈v4
    fotofreigabe = db.Column(db.Boolean, default=False)  # 👈 NEU v5
    mobil = db.Column(db.String(30), nullable=True)   # für WhatsApp, so wie eingegeben
    email = db.Column(db.String(200), nullable=True)
    hunde = db.relationship('Hund', backref='besitzer', lazy=True, cascade='all, delete-orphan')
    fotoeinwilligungen = db.relationship('Fotoeinwilligung', backref='person', lazy=True,
                                         cascade='all, delete-orphan',
                                         order_by='Fotoeinwilligung.hochgeladen_am.desc()')


    @property
    def vollstaendiger_name(self):
        return f"{self.vorname} {self.nachname}"

    @property
    def aktuelle_fotoeinwilligung(self):
        """Die neueste nicht widerrufene Einwilligung mit Dokument - oder None."""
        return next((e for e in self.fotoeinwilligungen if not e.widerrufen_am), None)

    @property
    def foto_status(self):
        """'nachgewiesen', 'ohne-nachweis' (nur angekreuzt), 'widerrufen' oder 'fehlt'."""
        if self.aktuelle_fotoeinwilligung:
            return 'nachgewiesen'
        if self.fotoeinwilligungen:
            return 'widerrufen'  # nach einem Widerruf nicht erneut nachfragen
        return 'ohne-nachweis' if self.fotofreigabe else 'fehlt'

    def _fotoeinwilligung_nachricht(self, kanal):
        """Bitte um die Einwilligung mit Link zum Formular - None, wenn nichts zu fragen ist."""
        if self.foto_status not in ('fehlt', 'ohne-nachweis'):
            return None
        return whatsapp.fotoeinwilligung_nachricht(
            self.vorname, [h.name for h in sorted(self.hunde, key=lambda h: h.name.lower())],
            einstellung('fotoeinwilligung_link'), kanal)

    @property
    def fotoeinwilligung_anfrage(self):
        """WhatsApp-Nachricht mit Link zum Formular - None, wenn nichts zu fragen ist oder die Nummer fehlt."""
        text = self._fotoeinwilligung_nachricht('whatsapp')
        return whatsapp.link(self.mobil, text) if text else None

    @property
    def fotoeinwilligung_mail(self):
        """Dieselbe Anfrage als E-Mail - None, wenn nichts zu fragen ist oder die Adresse fehlt."""
        text = self._fotoeinwilligung_nachricht('email')
        return mailto.link(self.email, 'Einwilligung zu Fotoaufnahmen', text) if text else None

    @property
    def kontaktdaten_fehlen(self):
        """Weder brauchbare Handynummer noch E-Mail - der Halter ist aus der App nicht erreichbar."""
        return not self.whatsapp_link and not self.email_link

    @property
    def kontaktdaten_anfrage(self):
        """WhatsApp an den Ansprechpartner für Kontaktdaten (Einstellungen) - None, wenn er nicht
        eingetragen ist oder der Halter erreichbar ist."""
        if not self.kontaktdaten_fehlen or not einstellung('kontaktdaten_name'):
            return None
        return whatsapp.link(einstellung('kontaktdaten_mobil'), whatsapp.kontaktdaten_nachricht(
            einstellung('kontaktdaten_name'), self.vollstaendiger_name,
            [h.name for h in sorted(self.hunde, key=lambda h: h.name.lower())]))

    @property
    def whatsapp_link(self):
        """Leerer Chat in WhatsApp Web - None ohne gültige Handynummer."""
        return whatsapp.link(self.mobil)

    @property
    def email_link(self):
        """Leere E-Mail an den Halter - None ohne gültige Adresse."""
        return mailto.link(self.email)

class Hund(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    geburtstag = db.Column(db.Date, nullable=True)
    gueltig_shp_dap_dhp = db.Column(db.Date, nullable=True)
    gueltig_l = db.Column(db.Date, nullable=True)
    gueltig_bbpi = db.Column(db.Date, nullable=True)
    gueltig_t = db.Column(db.Date, nullable=True)
    haftpflicht_gueltig = db.Column(db.Boolean, default=False)  # alte Ja/Nein-Angabe ohne Nachweis
    bemerkung = db.Column(db.Text, nullable=True)
    person_id = db.Column(db.Integer, db.ForeignKey('person.id'), nullable=False)
    nachweise = db.relationship('Nachweis', backref='hund', lazy=True,
                                cascade='all, delete-orphan',
                                order_by='Nachweis.hochgeladen_am.desc()')

    def nachweise_der_art(self, art):
        """Nachweise einer Art, der am längsten gültige zuerst (ohne Datum zuletzt)."""
        return sorted((n for n in self.nachweise if n.art == art),
                      key=lambda n: (n.gueltig_bis is not None, n.gueltig_bis or date.min, n.hochgeladen_am),
                      reverse=True)

    @property
    def aktueller_nachweis(self):
        """Der Haftpflicht-Nachweis mit der längsten Gültigkeit - ältere bleiben als Historie erhalten."""
        haftpflicht = self.nachweise_der_art('haftpflicht')
        return haftpflicht[0] if haftpflicht else None


# Alle Nachweise liegen in EINER Ablage - die Art bestimmt Felder und Ablauf.
# Neue Arten (z.B. Wesenstest) hier ergänzen; die Datei-Ablage und Sicherung bleiben gleich.
NACHWEIS_ARTEN = {
    'haftpflicht': {
        'name': 'Haftpflicht',
        'aussteller': 'Versicherer',
        'nummer': 'Vertragsnummer',
        'ausgestellt': 'Ausgestellt am',
        'gueltig_pflicht': True,     # ohne Gültigkeit kein Haftpflicht-Nachweis
        'tierdaten': True,           # Tier und Chip laut Dokument
        'impfungen': False,
    },
    'impfpass': {
        'name': 'Impfpass',
        'aussteller': 'Tierarzt / Praxis',
        'nummer': None,
        'ausgestellt': 'Letzte Impfung am (optional)',
        'gueltig_pflicht': False,
        'tierdaten': False,
        'impfungen': True,           # je Impfung "gültig bis" - wird beim Hund eingetragen
    },
}
# So lange vor Ablauf soll später an einen neuen Nachweis erinnert werden
ERINNERUNG_VORLAUF = relativedelta(weeks=6)


class Nachweis(db.Model):
    """Ein eingereichtes Dokument. Die Datei selbst liegt unveränderlich in instance/nachweise/."""
    id = db.Column(db.Integer, primary_key=True)
    hund_id = db.Column(db.Integer, db.ForeignKey('hund.id'), nullable=False)
    art = db.Column(db.String(20), nullable=False, default='haftpflicht')
    datei_sha256 = db.Column(db.String(64), nullable=False)
    datei_endung = db.Column(db.String(4), nullable=False)
    original_name = db.Column(db.String(255), nullable=True)
    hochgeladen_am = db.Column(db.DateTime, nullable=False, default=datetime.now)
    eingereicht_von = db.Column(db.String(200), nullable=True)   # Halter beim Hochladen
    aussteller = db.Column(db.String(200), nullable=True)        # Quelle: Versicherer, Tierarzt ...
    nummer = db.Column(db.String(100), nullable=True)            # Vertrags-/Dokumentnummer
    ausgestellt_am = db.Column(db.Date, nullable=True)
    gueltig_bis = db.Column(db.Date, nullable=True)
    tier_laut_nachweis = db.Column(db.String(100), nullable=True)
    chipnummer = db.Column(db.String(30), nullable=True)
    impfungen = db.Column(db.String(100), nullable=True)         # z.B. "SHP/DAP/DHP, L"
    # Impfpass: je Impfung "gültig bis" laut Pass, als JSON {"gueltig_l": "2027-12-05", ...}
    impf_gueltig = db.Column(db.Text, nullable=True)
    bemerkung = db.Column(db.Text, nullable=True)
    # Anzeige-Drehung der Datei in Grad (im Uhrzeigersinn) - die Datei selbst bleibt unverändert
    drehung = db.Column(db.Integer, nullable=True, default=0)
    # Weitere Fotos desselben Nachweises (z.B. mehrere Impfpass-Seiten); die erste Datei steht oben
    weitere_seiten = db.relationship('NachweisSeite', backref='nachweis', lazy=True,
                                     cascade='all, delete-orphan', order_by='NachweisSeite.nr')

    @property
    def seiten(self):
        """Alle Dateien des Nachweises in Reihenfolge - die erste ist der Nachweis selbst."""
        erste = {'sha256': self.datei_sha256, 'endung': self.datei_endung,
                 'original_name': self.original_name, 'drehung': self.drehung or 0}
        return [erste] + [{'sha256': s.datei_sha256, 'endung': s.datei_endung,
                           'original_name': s.original_name, 'drehung': s.drehung or 0}
                          for s in self.weitere_seiten]

    @property
    def impf_daten(self):
        """{feld: date} aus impf_gueltig - kaputte Einträge werden übergangen."""
        try:
            roh = json.loads(self.impf_gueltig or '{}')
        except ValueError:
            return {}
        daten = {}
        for feld, wert in roh.items():
            try:
                daten[feld] = date.fromisoformat(wert)
            except (TypeError, ValueError):
                continue
        return daten

    @property
    def art_name(self):
        return NACHWEIS_ARTEN.get(self.art, {}).get('name', self.art)

    @property
    def status(self):
        return get_status(self.gueltig_bis)

    @property
    def erinnern_ab(self):
        return self.gueltig_bis - ERINNERUNG_VORLAUF if self.gueltig_bis else None


class NachweisSeite(db.Model):
    """Zweite, dritte ... Datei eines Nachweises (z.B. Tollwut-Seite und Impfseite des Passes)."""
    id = db.Column(db.Integer, primary_key=True)
    nachweis_id = db.Column(db.Integer, db.ForeignKey('nachweis.id'), nullable=False)
    nr = db.Column(db.Integer, nullable=False)           # 2, 3, ... (1 ist der Nachweis selbst)
    datei_sha256 = db.Column(db.String(64), nullable=False)
    datei_endung = db.Column(db.String(4), nullable=False)
    original_name = db.Column(db.String(255), nullable=True)
    drehung = db.Column(db.Integer, nullable=True, default=0)


class Fotoeinwilligung(db.Model):
    """Unterschriebene Einwilligung zu Fotoaufnahmen - gilt für den Halter, nicht für einen Hund.
    Die Datei liegt wie alle Nachweise unveränderlich in instance/nachweise/."""
    id = db.Column(db.Integer, primary_key=True)
    person_id = db.Column(db.Integer, db.ForeignKey('person.id'), nullable=False)
    datei_sha256 = db.Column(db.String(64), nullable=False)
    datei_endung = db.Column(db.String(4), nullable=False)
    original_name = db.Column(db.String(255), nullable=True)
    hochgeladen_am = db.Column(db.DateTime, nullable=False, default=datetime.now)
    unterschrieben_am = db.Column(db.Date, nullable=True)
    widerrufen_am = db.Column(db.Date, nullable=True)   # Widerruf wird vermerkt, nie gelöscht
    bemerkung = db.Column(db.Text, nullable=True)


class Einstellung(db.Model):
    """Einstellungen, die Saskia selbst ändern kann (Einstellungsdialog oben rechts)."""
    schluessel = db.Column(db.String(50), primary_key=True)
    wert = db.Column(db.Text, nullable=True)


# Bekannte Einstellungen mit Standardwert - neue Einstellung = Eintrag hier
EINSTELLUNGEN = {
    'fotoeinwilligung_link': {
        'name': 'Link zum Formular „Einwilligung zu Fotoaufnahmen“',
        'hilfe': 'Wird in der WhatsApp-Anfrage mitgeschickt. Ändert der Verein das Formular, '
                 'hier den neuen Link eintragen.',
        'standard': 'https://www.hsv-grossbottwar.de/wp-content/uploads/2026/02/Einwilligung_Fotoaufnahmen_wolf.pdf',
        'art': 'url',
    },
    'kontaktdaten_name': {
        'name': 'Ansprechpartner für Kontaktdaten: Vor- und Nachname',
        'hilfe': 'Wer im Vorstand die Handynummern und E-Mail-Adressen der Mitglieder hat. '
                 'Fehlt bei einem Halter beides, gibt es einen Knopf „Kontaktdaten anfragen“.',
        'standard': '',
        'art': 'text',
    },
    'kontaktdaten_mobil': {
        'name': 'Ansprechpartner für Kontaktdaten: Handynummer',
        'hilfe': 'An diese Nummer geht die Anfrage per WhatsApp.',
        'standard': '',
        'art': 'tel',
        'platzhalter': 'z. B. 0171 1234567',
    },
    'theme': {
        'name': 'Farbschema',
        'hilfe': 'Wie der Hundemanager aussieht. Die Daten bleiben gleich.',
        'standard': 'tuerkis',
        'art': 'auswahl',
        'optionen': {'tuerkis': 'Türkis', 'gruen': 'Grün (bisher)'},
        # Farben der kleinen Vorschau-Karte im Einstellungsdialog
        'vorschau': {'tuerkis': {'kopf': '#1A928C', 'streifen': '#00CDCD', 'knopf': '#212121'},
                     'gruen': {'kopf': '#173B2C', 'streifen': '#173B2C', 'knopf': '#173B2C'}},
    },
}


# Die Hilfe "E-Mails mit GMX schreiben" steht neben den E-Mail-Knöpfen, bis Saskia so oft
# einen E-Mail-Link geklickt hat - danach nur noch unten im Fuß. Zähler in der Tabelle einstellung
# (nicht in EINSTELLUNGEN, er gehört nicht in den Einstellungsdialog).
MAIL_HILFE_BIS_KLICKS = 3
MAIL_KLICKS = 'mail_klicks'


def mail_klicks():
    eintrag = db.session.get(Einstellung, MAIL_KLICKS)
    try:
        return int(eintrag.wert) if eintrag and eintrag.wert else 0
    except ValueError:
        return 0


def einstellung(schluessel):
    """Gespeicherter Wert oder - wenn nie geändert - der Standard."""
    eintrag = db.session.get(Einstellung, schluessel)
    if eintrag and eintrag.wert:
        return eintrag.wert
    return EINSTELLUNGEN[schluessel]['standard']


def get_status(datum):
    """Prüft den Status eines Gültigkeitsdatums"""
    if datum is None:
        return 'keine-angabe'
    
    heute = date.today()
    ein_monat_spaeter = heute + relativedelta(months=1)
    
    if datum < heute:
        return 'abgelaufen'  # rot
    elif datum <= ein_monat_spaeter:
        return 'bald-ablaufend'  # orange
    else:
        return 'gueltig'  # grün

def format_datum(datum):
    """Formatiert ein Datum für die Anzeige"""
    if datum is None:
        return '-'
    return datum.strftime('%d.%m.%Y')

def iso_datum(wert):
    """Für <input type="date">: date-Objekt oder bereits 'JJJJ-MM-TT'."""
    if isinstance(wert, date):
        return wert.isoformat()
    return wert or ''

# Jinja2 Filter registrieren
app.jinja_env.filters['format_datum'] = format_datum
app.jinja_env.filters['iso_datum'] = iso_datum
app.jinja_env.globals['get_status'] = get_status

IMPFUNGEN = [
    ('SHP/DAP/DHP', 'gueltig_shp_dap_dhp'),
    ('L', 'gueltig_l'),
    ('BbPi', 'gueltig_bbpi'),
    ('T', 'gueltig_t'),
]
# Woran Saskia die Impfung im Impfpass erkennt (Aufkleber des Impfstoffs)
IMPF_ERKENNEN = {
    'gueltig_shp_dap_dhp': 'Staupe, Hepatitis, Parvovirose – Aufkleber z. B. Nobivac SHP, DHPPi, '
                           'Eurican DAPPi, Versican DHPPi',
    'gueltig_l': 'Leptospirose – Aufkleber z. B. Nobivac L4, Versican L4, Eurican L4',
    'gueltig_bbpi': 'Zwingerhusten – Aufkleber z. B. Nobivac BbPi / KC',
    'gueltig_t': 'Tollwut (Rabies) – eigene Seite im EU-Heimtierausweis, Aufkleber z. B. Nobivac T, '
                 'Rabisin, Rabikal',
}
# Ein "gültig bis" weiter in der Zukunft ist sicher ein Tippfehler (Tollwut: bis zu 3 Jahre)
IMPF_HOECHSTENS = relativedelta(years=5)
GESAMT_TEXT = {
    'ok': 'Alles gültig',
    'bald': 'Bald fällig',
    'abgelaufen': 'Abgelaufen',
    'unvollstaendig': 'Unvollständig',
}


def hund_ansicht(hund):
    """Alles, was eine Hundekarte braucht - inklusive Gesamtstatus für Farbe und Filter."""
    heute = date.today()
    impfungen = []
    for name, feld in IMPFUNGEN:
        datum = getattr(hund, feld)
        status = get_status(datum)
        hinweis = ''
        if status == 'abgelaufen':
            hinweis = 'abgelaufen'
        elif status == 'bald-ablaufend':
            tage = (datum - heute).days
            hinweis = 'heute' if tage == 0 else ('morgen' if tage == 1 else f'in {tage} Tagen')
        impfungen.append({'name': name, 'feld': feld, 'datum': datum, 'status': status, 'hinweis': hinweis,
                          'erkennen': IMPF_ERKENNEN[feld]})

    nachweis = hund.aktueller_nachweis
    if nachweis:
        haftpflicht = get_status(nachweis.gueltig_bis)
    elif hund.haftpflicht_gueltig:
        haftpflicht = 'ohne-nachweis'
    else:
        haftpflicht = 'keine-angabe'

    nachfragen = fehlende_angaben(hund, impfungen, haftpflicht, nachweis)
    whatsapp_nachfrage = email_nachfrage = None
    if nachfragen:
        whatsapp_nachfrage = whatsapp.link(hund.besitzer.mobil, whatsapp.fehlende_daten_nachricht(
            hund.besitzer.vorname, hund.name, nachfragen))
        email_nachfrage = mailto.link(hund.besitzer.email, f'{hund.name}: fehlende Angaben',
                                      whatsapp.fehlende_daten_nachricht(
                                          hund.besitzer.vorname, hund.name, nachfragen, 'email'))

    stati = [i['status'] for i in impfungen] + [haftpflicht]
    if 'abgelaufen' in stati:
        gesamt = 'abgelaufen'
    elif 'bald-ablaufend' in stati:
        gesamt = 'bald'
    elif 'keine-angabe' in stati or 'ohne-nachweis' in stati:
        gesamt = 'unvollstaendig'
    else:
        gesamt = 'ok'

    return {
        'hund': hund,
        'impfungen': impfungen,
        'nachweis': nachweis,
        'haftpflicht': haftpflicht,
        'gesamt': gesamt,
        'gesamt_text': GESAMT_TEXT[gesamt],
        'nachfragen': nachfragen,
        'whatsapp_nachfrage': whatsapp_nachfrage,
        'email_nachfrage': email_nachfrage,
        'impf_nachfrage': impf_nachfrage(hund, impfungen),
        'impf_mail': impf_nachfrage(hund, impfungen, 'email'),
    }


def impf_punkte(impfungen):
    """Fehlende, abgelaufene und bald fällige Impfungen - je Punkt eine Zeile für die Nachricht."""
    punkte = []
    for i in impfungen:
        if i['status'] == 'keine-angabe':
            punkte.append(f'Impfung {i["name"]}: Datum fehlt noch')
        elif i['status'] == 'abgelaufen':
            punkte.append(f'Impfung {i["name"]}: abgelaufen am {format_datum(i["datum"])}')
        elif i['status'] == 'bald-ablaufend':
            punkte.append(f'Impfung {i["name"]}: läuft am {format_datum(i["datum"])} ab')
    return punkte


def impf_nachfrage(hund, impfungen, kanal='whatsapp'):
    """WhatsApp- bzw. E-Mail-Link: Bitte um ein Foto der Impfpass-Seite - None, wenn alles
    gültig ist oder Handynummer bzw. Adresse fehlen."""
    punkte = impf_punkte(impfungen)
    if not punkte:
        return None
    text = whatsapp.impfpass_nachricht(hund.besitzer.vorname, hund.name, punkte, kanal)
    if kanal == 'email':
        return mailto.link(hund.besitzer.email, f'{hund.name}: Impfpass', text)
    return whatsapp.link(hund.besitzer.mobil, text)


def fehlende_angaben(hund, impfungen, haftpflicht, nachweis):
    """Was beim Halter nachzufragen ist - je Punkt eine kurze Zeile für die Nachricht."""
    punkte = impf_punkte(impfungen)
    if haftpflicht in ('keine-angabe', 'ohne-nachweis'):
        punkte.append('Haftpflicht: Versicherungsnachweis fehlt noch')
    elif haftpflicht == 'abgelaufen':
        punkte.append(f'Haftpflicht: Nachweis abgelaufen am {format_datum(nachweis.gueltig_bis)}')
    elif haftpflicht == 'bald-ablaufend':
        punkte.append(f'Haftpflicht: Nachweis läuft am {format_datum(nachweis.gueltig_bis)} ab')
    if not hund.geburtstag:
        punkte.append('Geburtstag')
    return punkte


# Filter über die Kennzahl-Kacheln: gelten wie die Kennzahlen nur für aktive Halter
KENNZAHL_FILTER = {
    'bald': lambda p, h: any(i['status'] == 'bald-ablaufend' for i in h['impfungen']),
    'abgelaufen': lambda p, h: any(i['status'] == 'abgelaufen' for i in h['impfungen']),
    'haftpflicht': lambda p, h: h['haftpflicht'] not in ('gueltig', 'bald-ablaufend'),
    'foto': lambda p, h: p.foto_status in ('fehlt', 'ohne-nachweis'),
}


@app.route('/')
def index():
    ansicht = request.args.get('ansicht', 'aktiv')
    if request.args.get('all') == '1':  # alte Links aus v5.1.1 und früher
        ansicht = 'alle'
    if ansicht not in ('aktiv', 'handlungsbedarf', 'alle'):
        ansicht = 'aktiv'
    filter_ = request.args.get('filter')
    if filter_ not in KENNZAHL_FILTER:
        filter_ = None
    if filter_:
        ansicht = 'aktiv'

    alle_personen = Person.query.order_by(
        Person.aktiv.desc(),
        func.lower(Person.nachname),
        func.lower(Person.vorname)
    ).all()
    aktive = [p for p in alle_personen if p.aktiv]
    pausierte = [p for p in alle_personen if not p.aktiv]

    gruppen = []
    for person in (alle_personen if ansicht == 'alle' else aktive):
        hunde = [hund_ansicht(h) for h in sorted(person.hunde, key=lambda h: h.name.lower())]
        if ansicht == 'handlungsbedarf':
            hunde = [h for h in hunde if h['gesamt'] != 'ok']
            if not hunde:
                continue
        if filter_:
            hunde = [h for h in hunde if KENNZAHL_FILTER[filter_](person, h)]
            if not hunde and not (filter_ == 'foto' and KENNZAHL_FILTER['foto'](person, None)):
                continue
        gruppen.append({'person': person, 'hunde': hunde})

    # Kennzahlen immer über die aktiven Halter
    aktive_hunde = [hund_ansicht(h) for p in aktive for h in p.hunde]
    impf_stati = [i['status'] for h in aktive_hunde for i in h['impfungen']]
    kennzahlen = {
        'hunde': len(aktive_hunde),
        'bald': impf_stati.count('bald-ablaufend'),
        'abgelaufen': impf_stati.count('abgelaufen'),
        'ohne_haftpflicht': sum(1 for h in aktive_hunde if h['haftpflicht'] not in ('gueltig', 'bald-ablaufend')),
        'handlungsbedarf': sum(1 for h in aktive_hunde if h['gesamt'] != 'ok'),
        'ohne_fotoeinwilligung': sum(1 for p in aktive if p.foto_status in ('fehlt', 'ohne-nachweis')),
    }

    return render_template('index.html', gruppen=gruppen, ansicht=ansicht, filter=filter_,
                           pausierte=pausierte if ansicht != 'alle' and not filter_ else [],
                           anzahl_pausiert=len(pausierte), kennzahlen=kennzahlen,
                           gibt_personen=bool(alle_personen))

def _person_angaben(person):
    """Übernimmt die Formularfelder. Gibt eine Fehlermeldung zurück oder None."""
    vorname = request.form.get('vorname', '').strip()
    nachname = request.form.get('nachname', '').strip()
    mobil = request.form.get('mobil', '').strip()[:30]
    email = request.form.get('email', '').strip()[:200]
    if not vorname or not nachname:
        return 'Bitte Vor- und Nachname angeben.'
    if mobil and not whatsapp.nummer(mobil):
        return 'Die Handynummer stimmt so nicht – bitte z. B. als 0171 1234567 oder +49 171 1234567 eingeben.'
    if email and not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', email):
        return 'Die E-Mail-Adresse stimmt so nicht – bitte prüfen.'
    person.vorname = vorname
    person.nachname = nachname
    # Mit unterschriebener Einwilligung bleibt die Freigabe bestehen (Häkchen ist dann gesperrt)
    person.fotofreigabe = request.form.get('fotofreigabe') == '1' or bool(
        person.id and person.aktuelle_fotoeinwilligung)
    person.mobil = mobil or None
    person.email = email or None
    return None


@app.route('/person/neu', methods=['GET', 'POST'])
def person_neu():
    if request.method == 'POST':
        person = Person()
        fehler = _person_angaben(person)
        if fehler:
            flash(fehler, 'error')
            return render_template('person_form.html', person=None, werte=request.form)

        db.session.add(person)
        db.session.commit()
        flash(f'Person "{person.vollstaendiger_name}" wurde angelegt.', 'success')
        return redirect(url_for('index'))
    
    return render_template('person_form.html', person=None)

@app.route('/person/<int:person_id>/bearbeiten', methods=['GET', 'POST'])
def person_bearbeiten(person_id):
    person = Person.query.get_or_404(person_id)
    
    if request.method == 'POST':
        fehler = _person_angaben(person)
        if fehler:
            flash(fehler, 'error')
            return render_template('person_form.html', person=person, werte=request.form)

        db.session.commit()
        flash(f'Person wurde aktualisiert.', 'success')
        return redirect(url_for('index'))
    
    return render_template('person_form.html', person=person)

@app.route('/person/<int:person_id>/loeschen', methods=['POST'])
def person_loeschen(person_id):
    person = Person.query.get_or_404(person_id)
    name = person.vollstaendiger_name
    db.session.delete(person)
    db.session.commit()
    flash(f'Person "{name}" und alle zugehörigen Hunde wurden gelöscht.', 'success')
    return redirect(url_for('index'))

@app.route('/person/<int:person_id>/hund/neu', methods=['GET', 'POST'])
def hund_neu(person_id):
    person = Person.query.get_or_404(person_id)
    
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        
        if not name:
            flash('Bitte mindestens den Namen des Hundes angeben.', 'error')
            return render_template('hund_form.html', person=person, hund=None)
        
        hund = Hund(
            name=name,
            geburtstag=parse_datum(request.form.get('geburtstag')),
            gueltig_shp_dap_dhp=parse_datum(request.form.get('gueltig_shp_dap_dhp')),
            gueltig_l=parse_datum(request.form.get('gueltig_l')),
            gueltig_bbpi=parse_datum(request.form.get('gueltig_bbpi')),
            gueltig_t=parse_datum(request.form.get('gueltig_t')),
            haftpflicht_gueltig=request.form.get('haftpflicht_gueltig') == 'ja',
            bemerkung=request.form.get('bemerkung', '').strip() or None,
            person_id=person.id
        )
        db.session.add(hund)
        db.session.commit()
        flash(f'Hund "{hund.name}" wurde angelegt.', 'success')
        return redirect(url_for('index'))
    
    return render_template('hund_form.html', person=person, hund=None)

@app.route('/hund/<int:hund_id>/bearbeiten', methods=['GET', 'POST'])
def hund_bearbeiten(hund_id):
    hund = Hund.query.get_or_404(hund_id)
    
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        
        if not name:
            flash('Bitte mindestens den Namen des Hundes angeben.', 'error')
            return render_template('hund_form.html', person=hund.besitzer, hund=hund)
        
        hund.name = name
        hund.geburtstag = parse_datum(request.form.get('geburtstag'))
        hund.gueltig_shp_dap_dhp = parse_datum(request.form.get('gueltig_shp_dap_dhp'))
        hund.gueltig_l = parse_datum(request.form.get('gueltig_l'))
        hund.gueltig_bbpi = parse_datum(request.form.get('gueltig_bbpi'))
        hund.gueltig_t = parse_datum(request.form.get('gueltig_t'))
        hund.haftpflicht_gueltig = request.form.get('haftpflicht_gueltig') == 'ja'
        hund.bemerkung = request.form.get('bemerkung', '').strip() or None
        
        db.session.commit()
        flash(f'Hund "{hund.name}" wurde aktualisiert.', 'success')
        return redirect(url_for('index'))
    
    return render_template('hund_form.html', person=hund.besitzer, hund=hund)

@app.route('/hund/<int:hund_id>/loeschen', methods=['POST'])
def hund_loeschen(hund_id):
    hund = Hund.query.get_or_404(hund_id)
    name = hund.name
    db.session.delete(hund)
    db.session.commit()
    flash(f'Hund "{name}" wurde gelöscht.', 'success')
    return redirect(url_for('index'))

# --- Nachweise (Haftpflicht, Impfpass, ...) ---

@app.errorhandler(413)
def datei_zu_gross(_fehler):
    flash('Die Datei ist zu groß (maximal 25 MB je Datei, 100 MB zusammen).', 'error')
    return redirect(request.referrer or url_for('index'))


@app.route('/hund/<int:hund_id>/haftpflicht')
def haftpflicht(hund_id):
    """Alte Adresse aus 5.1.1/5.1.2 - führt zur Nachweis-Seite des Hundes."""
    return redirect(url_for('hund_nachweise', hund_id=hund_id))


@app.route('/hund/<int:hund_id>/nachweise')
def hund_nachweise(hund_id):
    hund = Hund.query.get_or_404(hund_id)
    art = request.args.get('art', 'haftpflicht')
    return render_template('hund_nachweise.html', hund=hund, arten=NACHWEIS_ARTEN, ansicht=hund_ansicht(hund),
                           impf_namen=IMPFUNGEN, gewaehlte_art=art if art in NACHWEIS_ARTEN else 'haftpflicht')


MAX_DATEIEN = 10  # je Upload - ein Impfpass hat selten mehr beschriebene Seiten


def _angaben_seite(hund, nachweis, art, vorschlag, seiten, warnungen=(), text_erkannt=False):
    """Dialog zum Prüfen/Eintragen der Angaben - mit den Fotos daneben."""
    return render_template('nachweis_angaben.html', hund=hund, nachweis=nachweis, art=art,
                           art_info=NACHWEIS_ARTEN[art], impfungen_auswahl=[n for n, _ in IMPFUNGEN],
                           impfungen=hund_ansicht(hund)['impfungen'], vorschlag=vorschlag,
                           warnungen=list(warnungen), seiten=seiten, text_erkannt=text_erkannt,
                           heute=date.today().isoformat())


@app.route('/hund/<int:hund_id>/nachweise/hochladen', methods=['POST'])
@app.route('/hund/<int:hund_id>/haftpflicht/hochladen', methods=['POST'])
def nachweis_hochladen(hund_id):
    hund = Hund.query.get_or_404(hund_id)
    art = request.form.get('art', 'haftpflicht')
    if art not in NACHWEIS_ARTEN:
        art = 'haftpflicht'
    dateien = [d for d in request.files.getlist('datei') if d and d.filename]
    if not dateien:
        flash('Bitte eine Datei auswählen.', 'error')
        return redirect(url_for('hund_nachweise', hund_id=hund.id, art=art))
    if len(dateien) > MAX_DATEIEN:
        flash(f'Bitte höchstens {MAX_DATEIEN} Dateien auf einmal hochladen.', 'error')
        return redirect(url_for('hund_nachweise', hund_id=hund.id, art=art))

    seiten, erste_daten = [], None
    for datei in dateien:
        daten = datei.read()
        name = os.path.basename(datei.filename)[:255]
        try:
            sha256, endung = nachweise.speichern(daten)
        except nachweise.NachweisFehler as e:
            flash(f'{name}: {e}' if len(dateien) > 1 else str(e), 'error')
            return redirect(url_for('hund_nachweise', hund_id=hund.id, art=art))
        if any(s['sha256'] == sha256 for s in seiten):
            continue  # dieselbe Datei doppelt ausgewählt
        if erste_daten is None:
            erste_daten = daten
        # Quer oder auf dem Kopf fotografiert? Nur die Anzeige wird gedreht, nie die Datei
        drehung = 0 if endung == 'pdf' else bilder.ausrichtung_erkennen(daten)
        seiten.append({'sha256': sha256, 'endung': endung, 'original_name': name, 'drehung': drehung})

    vorschlag = {'eingereicht_von': hund.besitzer.vollstaendiger_name}
    warnungen = []
    erkannt = False
    if art == 'haftpflicht' and seiten[0]['endung'] == 'pdf':
        erkannt_daten = nachweise.angaben_vorschlagen(nachweise.pdf_text(erste_daten))
        erkannt = bool(erkannt_daten)
        vorschlag.update(erkannt_daten)
        tier = vorschlag.get('tier')
        if tier and hund.name.casefold() not in tier.casefold():
            warnungen.append(f'Im Nachweis steht das Tier „{tier}“ – passt das zu {hund.name}?')
    gueltig_bis = vorschlag.get('gueltig_bis')
    if gueltig_bis and gueltig_bis < date.today():
        warnungen.append('Laut Dokument ist dieser Nachweis bereits abgelaufen.')
    return _angaben_seite(hund, None, art, vorschlag, seiten, warnungen, erkannt)


def _seiten_aus_formular():
    """Die hochgeladenen Dateien aus den versteckten Feldern des Dialogs - jede muss abgelegt sein."""
    namen = request.form.getlist('original_name')
    drehungen = request.form.getlist('drehung')
    seiten = []
    for i, (sha256, endung) in enumerate(zip(request.form.getlist('sha256'), request.form.getlist('endung'))):
        if not os.path.exists(nachweise.pfad(sha256, endung)):
            raise nachweise.NachweisFehler('Die hochgeladene Datei wurde nicht gefunden.')
        seiten.append({'sha256': sha256, 'endung': endung,
                       'original_name': (namen[i] if i < len(namen) else '')[:255] or None,
                       'drehung': bilder.drehung_pruefen(drehungen[i] if i < len(drehungen) else 0)})
    if not seiten:
        raise nachweise.NachweisFehler('Die hochgeladene Datei wurde nicht gefunden.')
    return seiten


def _impf_angaben():
    """Je Impfung das "gültig bis" aus dem Formular -> ({feld: date}, Fehlermeldung oder None)."""
    daten = {}
    grenze = date.today() + IMPF_HOECHSTENS
    for name, feld in IMPFUNGEN:
        wert = request.form.get(f'impf_{feld}', '').strip()
        if not wert:
            continue
        datum = parse_datum(wert)
        if not datum or datum.year < 2000 or datum > grenze:
            return None, (f'Impfung {name}: „{format_datum(datum) if datum else wert}“ kann nicht stimmen '
                          f'– bitte das Datum im Impfpass noch einmal prüfen.')
        daten[feld] = datum
    return daten, None


def _nachweis_angaben(nachweis):
    """Übernimmt die Formularfelder. Gibt eine Fehlermeldung zurück oder None (dann ist nichts geändert)."""
    art_info = NACHWEIS_ARTEN[nachweis.art]
    gueltig_bis = parse_datum(request.form.get('gueltig_bis'))
    impf = {}
    if art_info['impfungen']:
        impf, fehler = _impf_angaben()
        if fehler:
            return fehler
        if impf:
            # Der Impfpass muss erneuert werden, sobald die erste Impfung abläuft
            gueltig_bis = min(impf.values())
    if art_info['gueltig_pflicht'] and not gueltig_bis:
        return 'Bitte angeben, bis wann der Nachweis gültig ist.'
    nachweis.gueltig_bis = gueltig_bis
    nachweis.ausgestellt_am = parse_datum(request.form.get('ausgestellt_am'))
    nachweis.eingereicht_von = request.form.get('eingereicht_von', '').strip()[:200] or None
    nachweis.aussteller = request.form.get('aussteller', '').strip()[:200] or None
    nachweis.nummer = request.form.get('nummer', '').strip()[:100] or None
    nachweis.tier_laut_nachweis = request.form.get('tier_laut_nachweis', '').strip()[:100] or None
    nachweis.chipnummer = request.form.get('chipnummer', '').strip()[:30] or None
    angehakt = request.form.getlist('impfungen')  # ältere Formulare: nur Häkchen, ohne Datum
    nachweis.impfungen = ', '.join(n for n, f in IMPFUNGEN if f in impf or n in angehakt) or None
    nachweis.impf_gueltig = json.dumps({f: d.isoformat() for f, d in impf.items()}) if impf else None
    nachweis.bemerkung = request.form.get('bemerkung', '').strip() or None
    return None


def impfungen_uebernehmen(hund, neu, vorher=None):
    """Trägt die Daten aus dem Impfpass beim Hund ein. Ein späteres Datum gewinnt: ein altes Foto
    überschreibt keine neuere Impfung. Ausnahme: Korrektur eines Nachweises (vorher = dessen alte
    Angaben) - was von ihm stammt, wird auch auf ein früheres Datum korrigiert.
    Gibt (übernommen, nicht übernommen) als lesbare Zeilen zurück."""
    vorher = vorher or {}
    uebernommen, behalten = [], []
    for name, feld in IMPFUNGEN:
        datum = neu.get(feld)
        if not datum:
            continue
        aktuell = getattr(hund, feld)
        if aktuell is None or datum >= aktuell or aktuell == vorher.get(feld):
            setattr(hund, feld, datum)
            uebernommen.append(f'{name} bis {format_datum(datum)}')
        else:
            behalten.append(f'{name} (eingetragen ist schon {format_datum(aktuell)})')
    return uebernommen, behalten


def _impf_meldung(uebernommen, behalten):
    meldung = ''
    if uebernommen:
        meldung += ' Beim Hund eingetragen: ' + ', '.join(uebernommen) + '.'
    if behalten:
        meldung += ' Nicht übernommen, weil schon ein späteres Datum eingetragen ist: ' + ', '.join(behalten) + '.'
    return meldung


def _formular_vorschlag():
    """Eingaben zurück ins Formular, wenn etwas nicht stimmt - nichts geht verloren."""
    vorschlag = request.form.to_dict()
    vorschlag['impfungen_liste'] = request.form.getlist('impfungen')
    return vorschlag


@app.route('/hund/<int:hund_id>/nachweise/speichern', methods=['POST'])
@app.route('/hund/<int:hund_id>/haftpflicht/speichern', methods=['POST'])
def nachweis_speichern(hund_id):
    hund = Hund.query.get_or_404(hund_id)
    art = request.form.get('art', 'haftpflicht')
    if art not in NACHWEIS_ARTEN:
        art = 'haftpflicht'
    try:
        seiten = _seiten_aus_formular()
    except nachweise.NachweisFehler as e:
        flash(f'{e} Bitte erneut hochladen.', 'error')
        return redirect(url_for('hund_nachweise', hund_id=hund.id, art=art))

    erste = seiten[0]
    nachweis = Nachweis(hund_id=hund.id, art=art, datei_sha256=erste['sha256'], datei_endung=erste['endung'],
                        original_name=erste['original_name'], drehung=erste['drehung'])
    for nr, seite in enumerate(seiten[1:], start=2):
        nachweis.weitere_seiten.append(NachweisSeite(
            nr=nr, datei_sha256=seite['sha256'], datei_endung=seite['endung'],
            original_name=seite['original_name'], drehung=seite['drehung']))
    fehler = _nachweis_angaben(nachweis)
    if fehler:
        flash(fehler, 'error')
        return _angaben_seite(hund, None, art, _formular_vorschlag(), seiten)
    db.session.add(nachweis)
    meldung = ''
    if NACHWEIS_ARTEN[art]['impfungen']:
        meldung = _impf_meldung(*impfungen_uebernehmen(hund, nachweis.impf_daten))
    db.session.commit()
    gueltig = f' (gültig bis {format_datum(nachweis.gueltig_bis)})' if nachweis.gueltig_bis else ''
    flash(f'{nachweis.art_name}-Nachweis für {hund.name} gespeichert{gueltig}.{meldung}', 'success')
    return redirect(url_for('hund_nachweise', hund_id=hund.id, art=art))


@app.route('/nachweis/<int:nachweis_id>/bearbeiten', methods=['GET', 'POST'])
def nachweis_bearbeiten(nachweis_id):
    """Angaben korrigieren und Fotos drehen - die Dateien selbst bleiben unverändert."""
    nachweis = Nachweis.query.get_or_404(nachweis_id)
    if request.method == 'POST':
        vorher = nachweis.impf_daten
        fehler = _nachweis_angaben(nachweis)
        if fehler:
            db.session.rollback()
            flash(fehler, 'error')
            seiten = nachweis.seiten
            for seite, wert in zip(seiten, request.form.getlist('drehung')):
                seite['drehung'] = bilder.drehung_pruefen(wert)
            return _angaben_seite(nachweis.hund, nachweis, nachweis.art, _formular_vorschlag(), seiten)
        for teil, wert in zip([nachweis] + list(nachweis.weitere_seiten), request.form.getlist('drehung')):
            teil.drehung = bilder.drehung_pruefen(wert)
        meldung = ''
        if NACHWEIS_ARTEN[nachweis.art]['impfungen']:
            meldung = _impf_meldung(*impfungen_uebernehmen(nachweis.hund, nachweis.impf_daten, vorher))
        db.session.commit()
        flash(f'Angaben zum Nachweis aktualisiert.{meldung}', 'success')
        return redirect(url_for('hund_nachweise', hund_id=nachweis.hund_id, art=nachweis.art))
    vorschlag = {
        'gueltig_bis': nachweis.gueltig_bis, 'ausgestellt_am': nachweis.ausgestellt_am,
        'eingereicht_von': nachweis.eingereicht_von, 'aussteller': nachweis.aussteller,
        'nummer': nachweis.nummer, 'tier': nachweis.tier_laut_nachweis,
        'chipnummer': nachweis.chipnummer, 'bemerkung': nachweis.bemerkung,
        'impfungen_liste': [i.strip() for i in (nachweis.impfungen or '').split(',') if i.strip()],
    }
    vorschlag.update({f'impf_{feld}': datum for feld, datum in nachweis.impf_daten.items()})
    return _angaben_seite(nachweis.hund, nachweis, nachweis.art, vorschlag, nachweis.seiten)


@app.route('/nachweis/<int:nachweis_id>/ansehen')
def nachweis_ansehen(nachweis_id):
    """Alle Seiten eines Nachweises richtig herum - mit Drehknöpfen und Link zum Original."""
    nachweis = Nachweis.query.get_or_404(nachweis_id)
    return render_template('nachweis_ansehen.html', nachweis=nachweis, hund=nachweis.hund)


@app.route('/nachweis/<int:nachweis_id>/drehen', methods=['POST'])
def nachweis_drehen(nachweis_id):
    """Eine Seite um 90 Grad weiterdrehen - gespeichert wird nur die Anzeige-Drehung."""
    nachweis = Nachweis.query.get_or_404(nachweis_id)
    teile = [nachweis] + list(nachweis.weitere_seiten)
    try:
        teil = teile[int(request.form.get('seite', 0))]
    except (ValueError, IndexError):
        return redirect(url_for('nachweis_ansehen', nachweis_id=nachweis.id))
    schritt = -90 if request.form.get('richtung') == 'links' else 90
    teil.drehung = bilder.drehung_pruefen((teil.drehung or 0) + schritt)
    db.session.commit()
    return redirect(url_for('nachweis_ansehen', nachweis_id=nachweis.id) + f'#seite-{teile.index(teil) + 1}')


@app.route('/nachweis/<int:nachweis_id>/datei')
def nachweis_datei(nachweis_id):
    nachweis = Nachweis.query.get_or_404(nachweis_id)
    return _sende_nachweis(nachweis.datei_sha256, nachweis.datei_endung, nachweis.original_name)


@app.route('/nachweis/vorschau/<sha256>.<endung>')
def nachweis_vorschau(sha256, endung):
    """Originaldatei ansehen, auch bevor die Angaben gespeichert sind."""
    return _sende_nachweis(sha256, endung, None)


@app.route('/nachweis/bild/<sha256>.<endung>')
def nachweis_bild(sha256, endung):
    """Foto richtig herum gedreht und in Bildschirmgröße. PDFs (und ohne Pillow) das Original."""
    try:
        datei_pfad = nachweise.pfad(sha256, endung)
    except nachweise.NachweisFehler:
        return 'Nicht gefunden', 404
    if endung != 'pdf' and os.path.exists(datei_pfad):
        with open(datei_pfad, 'rb') as f:
            ergebnis = bilder.ansicht(f.read(), request.args.get('drehung', 0))
        if ergebnis:
            antwort = send_file(io.BytesIO(ergebnis[0]), mimetype=ergebnis[1])
            # Inhalt und Drehung stecken in der Adresse - der Browser darf sich das Bild merken
            antwort.headers['Cache-Control'] = 'private, max-age=86400'
            return antwort
    return _sende_nachweis(sha256, endung, None)


def _sende_nachweis(sha256, endung, original_name):
    try:
        datei_pfad = nachweise.pfad(sha256, endung)
    except nachweise.NachweisFehler:
        return 'Nicht gefunden', 404
    if not os.path.exists(datei_pfad):
        return 'Die Datei fehlt - bitte unter "Sicherungen" einen Stand wiederherstellen.', 404
    return send_file(datei_pfad, mimetype=nachweise.mimetype(endung),
                     download_name=original_name or f'nachweis.{endung}')


def impfungen_faellig(personen, heute):
    """Hunde aktiver Halter mit Impfungen, die abgelaufen sind oder innerhalb von
    ERINNERUNG_VORLAUF ablaufen - früheste zuerst. Grundlage sind die Daten am Hund."""
    grenze = heute + ERINNERUNG_VORLAUF
    faellig = []
    for person in personen:
        if not person.aktiv:
            continue
        for hund in person.hunde:
            ansicht = hund_ansicht(hund)
            impfungen = [i for i in ansicht['impfungen'] if i['datum'] and i['datum'] <= grenze]
            if impfungen:
                faellig.append({'hund': hund, 'impfungen': impfungen,
                                'frueheste': min(i['datum'] for i in impfungen),
                                'whatsapp': ansicht['impf_nachfrage'],
                                'email': ansicht['impf_mail']})
    return sorted(faellig, key=lambda f: (f['frueheste'], f['hund'].name.lower()))


@app.route('/nachweise')
def nachweise_uebersicht():
    """Historie aller Nachweise - und welche bald erneuert werden müssen."""
    art = request.args.get('art', '')
    abfrage = Nachweis.query.join(Hund).join(Person)
    if art in NACHWEIS_ARTEN:
        abfrage = abfrage.filter(Nachweis.art == art)
    else:
        art = ''
    alle = abfrage.order_by(Nachweis.hochgeladen_am.desc()).all()

    # Bald fällig: nur der jeweils aktuelle Nachweis je Hund und Art zählt -
    # wer schon einen neueren eingereicht hat, muss nicht erinnert werden.
    # Impfungen zählen einzeln (am Hund) - dafür gibt es eine eigene Liste.
    heute = date.today()
    aktuellste = {}
    for n in alle:
        if n.gueltig_bis and n.hund.besitzer.aktiv and not NACHWEIS_ARTEN.get(n.art, {}).get('impfungen'):
            schluessel = (n.hund_id, n.art)
            if schluessel not in aktuellste or n.gueltig_bis > aktuellste[schluessel].gueltig_bis:
                aktuellste[schluessel] = n
    erneuern = sorted((n for n in aktuellste.values() if n.erinnern_ab <= heute),
                      key=lambda n: n.gueltig_bis)
    impf_faellig = impfungen_faellig(Person.query.all(), heute) if art in ('', 'impfpass') else []

    return render_template('nachweise.html', nachweise=alle, erneuern=erneuern, impf_faellig=impf_faellig,
                           art=art, arten=NACHWEIS_ARTEN, heute=heute, vorlauf_wochen=6)


# --- Fotoeinwilligung (je Halter) ---

@app.route('/person/<int:person_id>/fotoeinwilligung')
def fotoeinwilligung(person_id):
    person = Person.query.get_or_404(person_id)
    return render_template('fotoeinwilligung.html', person=person,
                           formular_link=einstellung('fotoeinwilligung_link'), heute=date.today())


@app.route('/person/<int:person_id>/fotoeinwilligung/hochladen', methods=['POST'])
def fotoeinwilligung_hochladen(person_id):
    person = Person.query.get_or_404(person_id)
    datei = request.files.get('datei')
    if not datei or not datei.filename:
        flash('Bitte die unterschriebene Einwilligung (PDF oder Foto) auswählen.', 'error')
        return redirect(url_for('fotoeinwilligung', person_id=person.id))
    unterschrieben_am = parse_datum(request.form.get('unterschrieben_am'))
    if unterschrieben_am and unterschrieben_am > date.today():
        flash('Das Datum der Unterschrift liegt in der Zukunft – bitte prüfen.', 'error')
        return redirect(url_for('fotoeinwilligung', person_id=person.id))
    try:
        sha256, endung = nachweise.speichern(datei.read())
    except nachweise.NachweisFehler as e:
        flash(str(e), 'error')
        return redirect(url_for('fotoeinwilligung', person_id=person.id))

    db.session.add(Fotoeinwilligung(
        person_id=person.id, datei_sha256=sha256, datei_endung=endung,
        original_name=os.path.basename(datei.filename)[:255], unterschrieben_am=unterschrieben_am,
        bemerkung=request.form.get('bemerkung', '').strip() or None))
    person.fotofreigabe = True
    db.session.commit()
    flash(f'Fotoeinwilligung von {person.vollstaendiger_name} gespeichert.', 'success')
    return redirect(url_for('fotoeinwilligung', person_id=person.id))


@app.route('/fotoeinwilligung/<int:einwilligung_id>/widerrufen', methods=['POST'])
def fotoeinwilligung_widerrufen(einwilligung_id):
    """Widerruf vermerken - Dokument und Eintrag bleiben als Nachweis erhalten."""
    einwilligung = Fotoeinwilligung.query.get_or_404(einwilligung_id)
    person = einwilligung.person
    einwilligung.widerrufen_am = parse_datum(request.form.get('widerrufen_am')) or date.today()
    person.fotofreigabe = person.aktuelle_fotoeinwilligung is not None
    db.session.commit()
    flash(f'Widerruf vom {format_datum(einwilligung.widerrufen_am)} vermerkt – '
          f'Fotos von {person.vollstaendiger_name} nicht mehr verwenden.', 'success')
    return redirect(url_for('fotoeinwilligung', person_id=person.id))


@app.route('/fotoeinwilligung/<int:einwilligung_id>/datei')
def fotoeinwilligung_datei(einwilligung_id):
    einwilligung = Fotoeinwilligung.query.get_or_404(einwilligung_id)
    return _sende_nachweis(einwilligung.datei_sha256, einwilligung.datei_endung, einwilligung.original_name)


# --- Einstellungen ---

def _einstellung_pruefen(schluessel, wert):
    """Fehlermeldung oder None."""
    if EINSTELLUNGEN[schluessel]['art'] == 'url' and wert and \
            not re.fullmatch(r'https?://[^\s/]+\.[^\s/]+(/\S*)?', wert):
        return f'{EINSTELLUNGEN[schluessel]["name"]}: bitte eine vollständige Adresse ' \
               f'eingeben, die mit https:// beginnt.'
    if EINSTELLUNGEN[schluessel]['art'] == 'tel' and wert and not whatsapp.nummer(wert):
        return f'{EINSTELLUNGEN[schluessel]["name"]}: stimmt so nicht – bitte z. B. als 0171 1234567 ' \
               f'oder +49 171 1234567 eingeben.'
    if EINSTELLUNGEN[schluessel]['art'] == 'auswahl' and wert and wert not in EINSTELLUNGEN[schluessel]['optionen']:
        return f'{EINSTELLUNGEN[schluessel]["name"]}: bitte einen Eintrag aus der Liste wählen.'
    return None


@app.template_filter('link_anzeige')
def link_anzeige(link):
    """Dateiname, Domain und Art für die Dokument-Karte im Einstellungsdialog - None ohne gültigen Link."""
    if not link or _einstellung_pruefen('fotoeinwilligung_link', link):
        return None
    teile = urllib.parse.urlsplit(link)
    name = urllib.parse.unquote(teile.path.rstrip('/').rsplit('/', 1)[-1])
    domain = teile.netloc[4:] if teile.netloc.startswith('www.') else teile.netloc
    return {'name': name or domain, 'domain': domain,
            'art': 'PDF' if name.lower().endswith('.pdf') else 'Link'}


@app.route('/einstellungen', methods=['GET', 'POST'])
def einstellungen():
    if request.method == 'POST':
        werte = {s: request.form.get(s, '').strip() for s in EINSTELLUNGEN}
        if request.form.get('standard'):
            werte = {s: '' for s in EINSTELLUNGEN}  # leer = Standard
        fehler = [f for f in (_einstellung_pruefen(s, w) for s, w in werte.items()) if f]
        if bool(werte['kontaktdaten_name']) != bool(werte['kontaktdaten_mobil']):
            fehler.append('Ansprechpartner für Kontaktdaten: bitte Name und Handynummer angeben (oder beides leer lassen).')
        if fehler:
            for f in fehler:
                flash(f, 'error')
            return render_template('einstellungen.html', werte=werte)
        for schluessel, wert in werte.items():
            eintrag = db.session.get(Einstellung, schluessel) or Einstellung(schluessel=schluessel)
            eintrag.wert = None if wert == EINSTELLUNGEN[schluessel]['standard'] else (wert or None)
            db.session.add(eintrag)
        db.session.commit()
        flash('Einstellungen gespeichert.', 'success')
        ziel = request.form.get('zurueck', '')
        # Nur Seiten dieser App als Rücksprung
        return redirect(ziel if ziel.startswith('/') and not ziel.startswith('//') else url_for('einstellungen'))
    return render_template('einstellungen.html', werte={s: einstellung(s) for s in EINSTELLUNGEN})


def parse_datum(datum_string):
    """Parst einen Datum-String in ein date-Objekt"""
    if not datum_string:
        return None
    try:
        return datetime.strptime(datum_string, '%Y-%m-%d').date()
    except ValueError:
        return None

@app.route('/export/excel')
def export_excel():
    """Exportiert alle Daten als Excel-Datei"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hundeübersicht"
    
    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    orange_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    fill_rot = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_gelb = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    fill_gruen = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Header
    headers = ['Nachname', 'Vorname', 'Foto', 'Hundename', 'Geburtstag', 'SHP/DAP/DHP', 'L', 'BbPi', 'T', 'Haftpflicht', 'Bemerkung', 'Handynummer', 'E-Mail']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    
    # Daten
    personen = Person.query.order_by(Person.nachname.asc(), Person.vorname.asc()).all()
    row = 2
    
    for person in personen:
        for hund in person.hunde:
            ws.cell(row=row, column=1, value=person.nachname).border = thin_border
            ws.cell(row=row, column=2, value=person.vorname).border = thin_border
            
            """ws.cell(row=row, column=3, value=person.fotofreigabe).border = thin_border"""
            foto_status = person.foto_status
            foto_cell = ws.cell(row=row, column=3, value={
                'nachgewiesen': 'Ja', 'ohne-nachweis': 'Ja (ohne Nachweis)',
                'widerrufen': 'Widerrufen'}.get(foto_status, 'Nein'))
            foto_cell.border = thin_border
            foto_cell.alignment = Alignment(horizontal='center')
            foto_cell.fill = {'nachgewiesen': green_fill, 'ohne-nachweis': orange_fill}.get(foto_status, red_fill)
                
                
            status = impfstatus_farbe(hund)
            ws.cell(row=row, column=4, value=hund.name).border = thin_border
            if status == 'rot':
                ws.cell(row=row, column=4, value=hund.name).fill = fill_rot
            elif status == 'gelb':
                ws.cell(row=row, column=4, value=hund.name).fill = fill_gelb
            elif status == 'gruen':
                ws.cell(row=row, column=4, value=hund.name).fill = fill_gruen

            
            ws.cell(row=row, column=5, value=format_datum(hund.geburtstag)).border = thin_border
            
            # Gültigkeitsdaten mit Farbcodierung
            for col, datum in [(6, hund.gueltig_shp_dap_dhp), (7, hund.gueltig_l), 
                               (8, hund.gueltig_bbpi), (9, hund.gueltig_t)]:
                cell = ws.cell(row=row, column=col, value=format_datum(datum))
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                status = get_status(datum)
                if status == 'gueltig':
                    cell.fill = green_fill
                elif status == 'bald-ablaufend':
                    cell.fill = orange_fill
                elif status == 'abgelaufen':
                    cell.fill = red_fill
            
            nachweis = hund.aktueller_nachweis
            if nachweis:
                # Mit Nachweis: Gültigkeitsdatum, gefärbt wie die Impfungen
                haftpflicht_cell = ws.cell(row=row, column=10, value=f'bis {format_datum(nachweis.gueltig_bis)}')
                haftpflicht_status = get_status(nachweis.gueltig_bis)
                haftpflicht_cell.fill = {'gueltig': green_fill, 'bald-ablaufend': orange_fill}.get(haftpflicht_status, red_fill)
            else:
                haftpflicht_cell = ws.cell(row=row, column=10,
                                           value='Ja (ohne Nachweis)' if hund.haftpflicht_gueltig else 'Nein')
                haftpflicht_cell.fill = orange_fill if hund.haftpflicht_gueltig else red_fill
            haftpflicht_cell.border = thin_border
            haftpflicht_cell.alignment = Alignment(horizontal='center')
            
            ws.cell(row=row, column=11, value=hund.bemerkung or '').border = thin_border
            ws.cell(row=row, column=12, value=person.mobil or '').border = thin_border
            ws.cell(row=row, column=13, value=person.email or '').border = thin_border
            row += 1
    
    # Spaltenbreiten anpassen
    column_widths = [15, 15, 18, 20, 12, 15, 12, 12, 12, 18, 30, 18, 28]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    
    # Als BytesIO speichern
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'hundeuebersicht_{date.today().strftime("%Y-%m-%d")}.xlsx'
    )



@app.route('/person/<int:person_id>/toggle_aktiv', methods=['POST'])
def person_toggle_aktiv(person_id):
    person = Person.query.get_or_404(person_id)
    person.aktiv = not person.aktiv
    db.session.commit()

    if person.aktiv:
        flash(f'{person.vollstaendiger_name} ist wieder aktiv.', 'success')
    else:
        flash(f'{person.vollstaendiger_name} ist pausiert – zu finden unten unter „Pausiert“.', 'success')

    return redirect(request.referrer or url_for('index'))


def impfstatus_farbe(hund):
    heute = date.today()
    plus_1_monat = heute + relativedelta(months=1)
    minus_1_monat = heute - relativedelta(months=1)

    impfungen = [
        hund.gueltig_shp_dap_dhp,
        hund.gueltig_l,
        hund.gueltig_bbpi,
        hund.gueltig_t
    ]

    # 🟥 ROT: fehlt oder >1 Monat abgelaufen
    for datum in impfungen:
        if datum is None:
            return 'rot'
        if datum < minus_1_monat:
            return 'rot'

    # 🟨 GELB: innerhalb ±1 Monat
    for datum in impfungen:
        if minus_1_monat <= datum <= plus_1_monat:
            return 'gelb'

    # 🟩 GRÜN: alles >1 Monat gültig
    return 'gruen'


# --- Updates ---

NEUSTART_CODE = 3        # start.py startet die App bei diesem Exit-Code neu
SELBSTTEST_FEHLER = 4    # start.py nimmt das Update dann zurück
# Beim Start festhalten - nach einem Update steht in VERSION schon die neue Nummer
LAUFENDE_VERSION = updater.aktuelle_version()
# Eindeutig je Prozess - daran erkennen Browser und Tests einen Neustart. Unter Windows
# ist der Server beim Neustart nicht erkennbar "weg": Verbindungsversuche warten dort
# ~2 Sekunden und landen dann nahtlos beim neuen Prozess.
INSTANZ = f'{os.getpid()}-{secrets.token_hex(4)}'
# Ab Protokoll 2 sichert start.py Updates ab (Rücknahme bei Fehlstart)
LAUNCHER_PROTOKOLL = int(os.environ.get('HUNDEMANAGER_LAUNCHER') or 0)

_hinweise = {'backup_fehler': None}


def lies_update_fehlgeschlagen():
    try:
        with open(backup.UPDATE_FEHLGESCHLAGEN, encoding='utf-8') as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


@app.context_processor
def update_kontext():
    return {
        'app_version': LAUFENDE_VERSION,
        'update_info': updater.update_status(),
        'update_fehlgeschlagen': lies_update_fehlgeschlagen(),
        'backup_fehler': _hinweise['backup_fehler'],
        # Alter Starter (v5.1.0) läuft noch - erst nach Neustart über START.bat ist die Absicherung aktiv
        'starter_veraltet': LAUNCHER_PROTOKOLL == 1,
    }


KONTAKTDATEN_VORNAME_MARKE = '\x00vorname\x00'


@app.context_processor
def einstellungen_kontext():
    # Für den Einstellungsdialog im Kopf jeder Seite
    return {'einstellungen_info': EINSTELLUNGEN,
            'einstellungen_werte': {s: einstellung(s) for s in EINSTELLUNGEN},
            # Beispiel der Anfrage an den Ansprechpartner - der Vorname wird im Dialog live eingesetzt
            'kontaktdaten_beispiel': whatsapp.kontaktdaten_nachricht(
                KONTAKTDATEN_VORNAME_MARKE, 'Max Mustermann', ['Bello']),
            'kontaktdaten_vorname_marke': KONTAKTDATEN_VORNAME_MARKE}


@app.context_processor
def mail_hilfe_kontext():
    return {'mail_hilfe_zeigen': mail_klicks() < MAIL_HILFE_BIS_KLICKS}


@app.route('/email/geklickt', methods=['POST'])
def email_geklickt():
    """Ein E-Mail-Link wurde geklickt (meldet base.html) - zählt bis zum Ausblenden der Hilfe."""
    klicks = mail_klicks()
    if klicks < MAIL_HILFE_BIS_KLICKS:
        eintrag = db.session.get(Einstellung, MAIL_KLICKS) or Einstellung(schluessel=MAIL_KLICKS)
        eintrag.wert = str(klicks + 1)
        db.session.add(eintrag)
        db.session.commit()
    return '', 204


@app.route('/update/pruefen', methods=['POST'])
def update_pruefen():
    info = updater.pruefe_auf_update(erzwingen=True)
    if info.get('fehler'):
        flash(info['fehler'] + ' - bitte Internetverbindung prüfen.', 'error')
    elif info['verfuegbar']:
        flash(f'Version {info["version"]} ist verfügbar.', 'success')
    else:
        flash('Der Hundemanager ist auf dem neuesten Stand.', 'success')
    return redirect(url_for('index'))


@app.route('/update/kopf')
def update_kopf():
    """Nur der Update-Knopf - offene Seiten fragen ihn jede Minute nach."""
    return render_template('_update_kopf.html')


@app.route('/update/installieren', methods=['POST'])
def update_installieren():
    # Vorher laden: nach dem Update liegen auf der Platte schon die neuen Templates,
    # dieser Prozess kennt aber nur die alten Routen
    fertig_seite = app.jinja_env.get_template('update_fertig.html')
    erfolg, meldung = updater.installiere_update()
    if not erfolg:
        flash(meldung, 'error')
        return redirect(url_for('index'))

    threading.Thread(target=neustart, daemon=True).start()
    return fertig_seite.render(meldung=meldung, alte_version=LAUFENDE_VERSION, alte_instanz=INSTANZ,
                               url_for=url_for)


@app.route('/update/hinweis-schliessen', methods=['POST'])
def update_hinweis_schliessen():
    if os.path.exists(backup.UPDATE_FEHLGESCHLAGEN):
        os.remove(backup.UPDATE_FEHLGESCHLAGEN)
    return redirect(url_for('index'))


@app.route('/api/version')
def api_version():
    return jsonify(version=LAUFENDE_VERSION, instanz=INSTANZ)


def neustart():
    time.sleep(1.5)  # Antwortseite erst noch ausliefern
    if LAUNCHER_PROTOKOLL:
        os._exit(NEUSTART_CODE)  # start.py startet neu
    # Ohne start.py gestartet (z.B. "python app.py"): selbst neu starten
    env = dict(os.environ, HUNDEMANAGER_WARTEN='2')
    subprocess.Popen([sys.executable] + sys.argv, env=env)
    os._exit(0)


# --- Sicherungen ---

def taeglich_sichern():
    """Einmal pro Tag eine Sicherung - schützt auch vor versehentlichem Löschen."""
    try:
        if backup.taegliches_backup_faellig():
            _, spiegel_fehler = backup.erstelle_backup('taeglich', LAUFENDE_VERSION)
            _hinweise['backup_fehler'] = spiegel_fehler
    except backup.BackupFehler as e:
        _hinweise['backup_fehler'] = f'Tägliche Sicherung fehlgeschlagen: {e}'
        print(_hinweise['backup_fehler'])


@app.route('/sicherungen')
def sicherungen():
    return render_template('sicherungen.html', backups=backup.liste_backups(),
                           backup_dir=backup.BACKUP_DIR, zweiter_ort=backup.zweiter_ort())


@app.route('/sicherungen/jetzt', methods=['POST'])
def sicherung_jetzt():
    try:
        _, spiegel_fehler = backup.erstelle_backup('manuell', LAUFENDE_VERSION)
        flash('Sicherung erstellt und geprüft.', 'success')
        if spiegel_fehler:
            flash(spiegel_fehler, 'error')
    except backup.BackupFehler as e:
        flash(f'Sicherung fehlgeschlagen: {e}', 'error')
    return redirect(url_for('sicherungen'))


@app.route('/sicherungen/wiederherstellen', methods=['POST'])
def sicherung_wiederherstellen():
    try:
        ordner = backup.finde_backup(request.form.get('name', ''))
        # Keine offenen Verbindungen der App, während die Datenbank zurückgespielt wird
        db.session.remove()
        db.engine.dispose()
        manifest = backup.datenbank_wiederherstellen(ordner, LAUFENDE_VERSION)
        migriere_datenbank()  # ältere Sicherung? Fehlende Spalten ergänzen
        flash(f'Daten vom {manifest["erstellt"][:16].replace("T", " ")} wiederhergestellt. '
              f'Der vorherige Stand wurde vorher gesichert.', 'success')
    except backup.BackupFehler as e:
        flash(f'Wiederherstellung nicht möglich: {e} Die Daten wurden nicht verändert.', 'error')
    return redirect(url_for('sicherungen'))


# --- Datenbank-Migration ---

def _sql_standardwert(spalte):
    if spalte.default is None or not spalte.default.is_scalar:
        return None
    wert = spalte.default.arg
    if isinstance(wert, bool):
        return '1' if wert else '0'
    if isinstance(wert, (int, float)):
        return str(wert)
    return "'" + str(wert).replace("'", "''") + "'"


def migriere_datenbank():
    """Ergänzt fehlende Tabellen/Spalten in einer bestehenden Datenbank.

    db.create_all() legt nur neue Tabellen an, aber keine neuen Spalten.
    Ohne diese Migration würde eine ältere Datenbank (z.B. aus v3/v4) nach
    einem Update mit "no such column" abstürzen.
    """
    inspektor = inspect(db.engine)
    vorhandene_tabellen = set(inspektor.get_table_names())
    gesichert = False

    for tabelle in db.metadata.sorted_tables:
        if tabelle.name not in vorhandene_tabellen:
            continue  # wird unten von create_all() angelegt
        vorhanden = {s['name'] for s in inspektor.get_columns(tabelle.name)}
        for spalte in tabelle.columns:
            if spalte.name in vorhanden:
                continue
            if not gesichert:
                # Ohne geprüfte Sicherung keine Änderung an der Datenbank
                backup.erstelle_backup('vor-migration', LAUFENDE_VERSION, mit_code=False)
                gesichert = True
            sql = f'ALTER TABLE {tabelle.name} ADD COLUMN {spalte.name} {spalte.type.compile(db.engine.dialect)}'
            standard = _sql_standardwert(spalte)
            if standard is not None:
                sql += f' DEFAULT {standard}'
            with db.engine.begin() as verbindung:
                verbindung.execute(text(sql))
            print(f'Datenbank aktualisiert: Spalte {tabelle.name}.{spalte.name} ergänzt')

    db.create_all()
    _haftpflicht_nachweise_uebernehmen(gesichert)


def _haftpflicht_nachweise_uebernehmen(schon_gesichert):
    """5.1.1/5.1.2 hatten eine eigene Tabelle haftpflicht_nachweis - in die gemeinsame
    Ablage 'nachweis' übernehmen. Die alte Tabelle bleibt zur Sicherheit unverändert stehen."""
    if 'haftpflicht_nachweis' not in inspect(db.engine).get_table_names():
        return
    with db.engine.connect() as v:
        alt = v.execute(text('SELECT COUNT(*) FROM haftpflicht_nachweis')).scalar()
        schon = v.execute(text("SELECT COUNT(*) FROM nachweis WHERE art = 'haftpflicht'")).scalar()
    if alt == 0 or schon > 0:
        return  # nichts zu tun oder bereits übernommen
    if not schon_gesichert:
        backup.erstelle_backup('vor-migration', LAUFENDE_VERSION, mit_code=False)
    with db.engine.begin() as v:  # alles oder nichts
        v.execute(text('''
            INSERT INTO nachweis (hund_id, art, datei_sha256, datei_endung, original_name,
                                  hochgeladen_am, eingereicht_von, aussteller, nummer,
                                  ausgestellt_am, gueltig_bis, tier_laut_nachweis, chipnummer)
            SELECT h.hund_id, 'haftpflicht', h.datei_sha256, h.datei_endung, h.original_name,
                   h.hochgeladen_am, p.vorname || ' ' || p.nachname, h.versicherer, h.vertragsnummer,
                   h.ausgestellt_am, h.gueltig_bis, h.tier_laut_nachweis, h.chipnummer
            FROM haftpflicht_nachweis h
            JOIN hund ON hund.id = h.hund_id
            JOIN person p ON p.id = hund.person_id
            ORDER BY h.id'''))
        neu = v.execute(text("SELECT COUNT(*) FROM nachweis WHERE art = 'haftpflicht'")).scalar()
        if neu != alt:
            raise RuntimeError(f'Übernahme der Haftpflicht-Nachweise unvollständig ({neu} von {alt})')
    print(f'Datenbank aktualisiert: {alt} Haftpflicht-Nachweis(e) in die Nachweis-Ablage übernommen')


def selbsttest():
    """Lädt die wichtigsten Seiten einmal intern - fällt etwas um, gilt das Update als fehlgeschlagen."""
    with app.test_client() as client:
        for pfad in ('/', '/?ansicht=handlungsbedarf', '/?ansicht=alle', '/nachweise', '/sicherungen', '/einstellungen',
                     '/export/excel'):
            antwort = client.get(pfad)
            if antwort.status_code != 200:
                raise RuntimeError(f'Selbsttest: {pfad} liefert {antwort.status_code}')


def starten():
    nach_update = backup.lies_update_marker() is not None
    try:
        with app.app_context():
            migriere_datenbank()
            selbsttest()
    except Exception:
        traceback.print_exc()
        if not nach_update:
            raise
        if LAUNCHER_PROTOKOLL >= 2:
            sys.exit(SELBSTTEST_FEHLER)  # start.py nimmt das Update zurück
        # Alter Starter: selbst zurücknehmen und neu starten lassen
        backup.update_zuruecknehmen('Die neue Version hat den Selbsttest nicht bestanden.')
        neustart()

    if nach_update:
        backup.entferne_update_marker()
        print(f'Update auf {LAUFENDE_VERSION} bestätigt.')

    taeglich_sichern()
    updater.starte_hintergrund_pruefung(taeglich_sichern)
    port = int(os.environ.get('HUNDEMANAGER_PORT', 5000))
    print(f'Hundemanager {LAUFENDE_VERSION} läuft auf http://127.0.0.1:{port}')
    # Nur auf diesem PC erreichbar, ohne Debugger
    app.run(host='127.0.0.1', port=port)


if __name__ == '__main__':
    time.sleep(float(os.environ.get('HUNDEMANAGER_WARTEN', 0)))
    starten()
